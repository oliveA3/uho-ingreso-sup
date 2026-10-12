from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook
from rest_framework.test import APIRequestFactory, APITestCase
from django.contrib.auth.models import AnonymousUser
from django.contrib.auth import get_user_model
from django.utils import timezone

from apps.authentication.models import Estudiante
from apps.gestion_personal.models import BoletaInteres, BoletaInteresItem, BoletaSolicitud, BoletaSolicitudItem, ResultadoExamen
from apps.gestion_provincial.models import ETAPAS_NOMBRES, CorteCarrera, Etapa, Otorgamiento, PlanPlaza, Proceso
from apps.import_export.services.carreras_service import CareerExcelService
from apps.superadmin.models import Asignatura, Carrera, Ces, Escuela, IdentidadVisual, Municipio, Provincia, TipoOtorgamiento


class PublicLandingCortesTests(APITestCase):
    def setUp(self):
        self.province = Provincia.objects.create(nombre="Provincia de cortes")
        self.other_province = Provincia.objects.create(nombre="Otra provincia de cortes")
        ces = Ces.objects.create(nombre="CES de cortes")
        award_type, _ = TipoOtorgamiento.objects.get_or_create(nombre="Provincial")
        self.career = Carrera.objects.create(
            codigo="COR-1", nombre="Carrera con corte", ces=ces, provincia=self.province,
        )
        other_career = Carrera.objects.create(
            codigo="COR-2", nombre="Carrera de otra provincia", ces=ces, provincia=self.other_province,
        )
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[6])
        process = Proceso.objects.create(anio=date(2025, 1, 1), etapa=stage)
        PlanPlaza.objects.create(
            proceso=process, carrera=self.career, cantidad_plazas=1,
            otorgamiento_tipo=award_type, ces=ces, provincia=self.province, sexo="A",
        )
        PlanPlaza.objects.create(
            proceso=process, carrera=other_career, cantidad_plazas=1,
            otorgamiento_tipo=award_type, ces=ces, provincia=self.other_province, sexo="A",
        )
        CorteCarrera.objects.create(proceso=process, carrera=self.career, indice_corte=92.5)
        CorteCarrera.objects.create(proceso=process, carrera=other_career, indice_corte=88.0)

    def test_anonymous_user_can_read_historical_cuts(self):
        from .views import LandingCortesView

        request = APIRequestFactory().get("/api/import-export/cortes/landing/", {"anio": 2025})
        view = LandingCortesView()
        drf_request = view.initialize_request(request)
        drf_request.user = AnonymousUser()
        response = view.get(drf_request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["year"], 2025)
        self.assertEqual(len(response.data["items"]), 2)

    def test_anonymous_user_can_download_filtered_historical_cuts(self):
        from .views import LandingExcelExportView

        request = APIRequestFactory().get(
            "/api/import-export/landing/export/cortes/",
            {"anio": 2025, "provincia": self.province.nombre},
        )
        request.user = AnonymousUser()
        view = LandingExcelExportView()
        response = view.get(view.initialize_request(request), "cortes")

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.values)
        self.assertEqual(rows[1][0], "COR-1")
        self.assertEqual(len(rows), 2)

    def test_pdf_builder_includes_branding_header(self):
        IdentidadVisual.objects.create(
            nombre_sistema="Sistema de Prueba",
            logo_url="https://example.com/logo.png",
        )

        from .views import build_pdf_document
        from apps.core.pdf import build_report_html

        pdf_bytes = build_pdf_document(
            title="Boleta de Prueba",
            lines=["Línea 1", "Línea 2"],
            process_name="Proceso de Ingreso 2026",
            process_identifier="Etapa 3 - Plan de plazas",
        )

        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        content = build_report_html(
            "Boleta de Prueba",
            ["Línea 1", "Línea 2"],
            ["Proceso: Proceso de Ingreso 2026", "Identificación del proceso: Etapa 3 - Plan de plazas"],
        )
        self.assertIn("Sistema de Prueba", content)
        self.assertIn("Proceso de Ingreso 2026", content)
        self.assertIn("Etapa 3 - Plan de plazas", content)
        self.assertIn("Fecha de generación", content)


class PublicResultsLandingTests(APITestCase):
    def setUp(self):
        self.province = Provincia.objects.create(nombre="Provincia protegida")
        municipality = Municipio.objects.create(nombre="Municipio protegido", provincia=self.province)
        school = Escuela.objects.create(nombre="Escuela protegida", municipio=municipality)
        ces = Ces.objects.create(nombre="CES protegido")
        self.career = Carrera.objects.create(
            codigo="PRV-1", nombre="Carrera protegida", ces=ces, provincia=self.province,
        )
        self.rare_career = Carrera.objects.create(
            codigo="PRV-2", nombre="Carrera minoritaria", ces=ces, provincia=self.province,
        )
        self.math = Asignatura.objects.create(nombre="Asignatura protegida")
        self.rare_subject = Asignatura.objects.create(nombre="Asignatura minoritaria")
        self.results_process = Proceso.objects.create(
            anio=date(2025, 1, 1),
            etapa=Etapa.objects.get(nombre=ETAPAS_NOMBRES[5]),
        )
        self.awards_process = Proceso.objects.create(
            anio=date(2025, 1, 1),
            etapa=Etapa.objects.get(nombre=ETAPAS_NOMBRES[6]),
        )
        students = []
        for index in range(6):
            student = Estudiante.objects.create(
                ci=f"900000000{index:02d}",
                nombre=f"Nombre privado {index}",
                apellidos=f"Apellido privado {index}",
                sexo="F",
                direccion="Dirección privada",
                escuela=school,
            )
            students.append(student)
            if index < 5:
                ResultadoExamen.objects.create(
                    estudiante=student,
                    proceso=self.results_process,
                    asignatura=self.math,
                    nota=80 + index,
                    fecha_limite_reclamo=date(2025, 12, 31),
                )
                Otorgamiento.objects.create(
                    estudiante=student,
                    proceso=self.awards_process,
                    carrera=self.career,
                    indice_otorgamiento=90 + index,
                )
        ResultadoExamen.objects.create(
            estudiante=students[5],
            proceso=self.results_process,
            asignatura=self.rare_subject,
            nota=99,
            fecha_limite_reclamo=date(2025, 12, 31),
        )
        Otorgamiento.objects.create(
            estudiante=students[5],
            proceso=self.awards_process,
            carrera=self.rare_career,
            indice_otorgamiento=99,
        )
        self.factory = APIRequestFactory()

    def test_public_results_include_student_rows_and_apply_filters(self):
        from .views import LandingResultadosView

        response = LandingResultadosView.as_view()(
            self.factory.get(
                "/api/v1/import-export/resultados/landing/",
                {"anio": 2025, "municipio": "Municipio protegido", "escuela": "Escuela protegida"},
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["results"]), 6)
        result = response.data["results"][0]
        self.assertEqual(result["student"], "Nombre privado 0 Apellido privado 0")
        self.assertEqual(result["school"], "Escuela protegida")
        self.assertEqual(result["municipality"], "Municipio protegido")
        self.assertEqual(result["province"], "Provincia protegida")
        self.assertEqual(len(response.data["municipalities"]), 1)
        self.assertEqual(len(response.data["schools"]), 1)

    def test_public_awards_include_student_rows_and_filter_by_ci(self):
        from .views import LandingOtorgamientosView

        response = LandingOtorgamientosView.as_view()(
            self.factory.get(
                "/api/v1/import-export/otorgamiento/landing/",
                {"anio": 2025, "ci": "90000000000"},
            )
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["results"]), 1)
        award = response.data["results"][0]
        self.assertEqual(award["student"], "Nombre privado 0 Apellido privado 0")
        self.assertEqual(award["ci"], "90000000000")
        self.assertEqual(award["school"], "Escuela protegida")

    def test_public_excel_exports_include_student_results_and_awards(self):
        from .views import LandingExcelExportView

        for kind in ("resultados", "otorgamientos"):
            with self.subTest(kind=kind):
                response = LandingExcelExportView.as_view()(
                    self.factory.get(f"/api/v1/import-export/landing/export/{kind}/", {"anio": 2025}),
                    kind=kind,
                )
                workbook = load_workbook(BytesIO(response.content), read_only=True)
                rows = list(workbook.active.values)
                values = [str(value) for row in rows for value in row]
                self.assertIn("CI", rows[0])
                self.assertIn("Estudiante", rows[0])
                self.assertTrue(any("privado" in value.lower() for value in values))
                self.assertTrue(any(value.startswith("900000000") for value in values))


class StudentBallotExcelExportTests(APITestCase):
    def setUp(self):
        province = Provincia.objects.create(nombre="Provincia de boletas")
        municipality = Municipio.objects.create(nombre="Municipio de boletas", provincia=province)
        school = Escuela.objects.create(nombre="Escuela de boletas", municipio=municipality)
        ces = Ces.objects.create(nombre="CES de boletas")
        career = Carrera.objects.create(
            codigo="BOL-1", nombre="Carrera de boletas", ces=ces, provincia=province,
        )
        award_type = TipoOtorgamiento.objects.create(nombre="Tipo de boletas")
        stage_interest = Etapa.objects.get_or_create(nombre=ETAPAS_NOMBRES[2])[0]
        stage_solicitud = Etapa.objects.get_or_create(nombre=ETAPAS_NOMBRES[3])[0]
        interest_process = Proceso.objects.create(anio=timezone.now().date().replace(month=1, day=1), etapa=stage_interest)
        solicitud_process = Proceso.objects.create(anio=timezone.now().date().replace(month=1, day=1), etapa=stage_solicitud)
        plan = PlanPlaza.objects.create(
            proceso=solicitud_process, carrera=career, cantidad_plazas=10,
            otorgamiento_tipo=award_type, ces=ces, provincia=province, sexo="A",
        )
        user = get_user_model().objects.create_user(
            username="ballot.export", email="ballot.export@example.com", password="secret1234", rol="estudiante",
        )
        student = Estudiante.objects.create(
            usuario=user, ci="12345678901", nombre="Ada", apellidos="Lovelace", sexo="F",
            direccion="Dirección de prueba", escuela=school,
        )
        interest = BoletaInteres.objects.create(estudiante=student, proceso=interest_process)
        BoletaInteresItem.objects.create(boleta_interes=interest, carrera=career, prioridad=1)
        solicitud = BoletaSolicitud.objects.create(estudiante=student, proceso=solicitud_process)
        BoletaSolicitudItem.objects.create(boleta_solicitud=solicitud, plan_plaza=plan, prioridad=1)
        self.user = user

    def test_student_ballots_can_be_exported_as_excel(self):
        for endpoint, filename, expected_title in [
            ("/api/v1/import-export/export/boleta-interes/excel/", "boleta-interes.xlsx", "BOLETA DE INTERES"),
            ("/api/v1/import-export/export/boleta-solicitud/excel/", "boleta-solicitud.xlsx", "BOLETA DE SOLICITUD"),
        ]:
            with self.subTest(endpoint=endpoint):
                from .views import StudentInterestExcelExportView, SolicitudExcelExportView

                view_class = StudentInterestExcelExportView if "interes" in endpoint else SolicitudExcelExportView
                request = APIRequestFactory().get(endpoint)
                view = view_class()
                drf_request = view.initialize_request(request)
                drf_request.user = self.user
                response = view.get(drf_request)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.assertIn(filename, response["Content-Disposition"])
                rows = list(load_workbook(BytesIO(response.content), read_only=True).active.values)
                self.assertEqual(rows[0][0], expected_title)
                self.assertEqual(rows[6][0], "Prioridad")
                self.assertEqual(rows[7][2], "Carrera de boletas")

class EscalafonRankingTests(APITestCase):
    def test_rank_orders_by_index_then_names_and_assigns_position(self):
        from types import SimpleNamespace as N
        from apps.import_export.services.escalafon_service import rank_escalafon_entries
        def e(i, idx, ap, no):
            return N(id=i, indice_general=idx, estudiante=N(apellidos=ap, nombre=no))
        entries = [e(1, 90, "Perez", "Ana"), e(2, 95, "Zeta", "Luis"), e(3, 90, "Alonso", "Zoe"), e(4, 90, "Perez", "Ana")]
        ranked = rank_escalafon_entries(entries)
        self.assertEqual([(p, x.id) for p, x in ranked], [(1, 2), (2, 3), (3, 1), (4, 4)])


class ImportTaskStatusTests(APITestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user(
            username="task.owner", email="task.owner@example.com", password="secret1234", rol="jefe_comision",
        )
        self.other = get_user_model().objects.create_user(
            username="task.other", email="task.other@example.com", password="secret1234", rol="jefe_comision",
        )

    def _upload(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Codigo", "Nombre", "Ces", "Provincia"])
        buffer = BytesIO()
        workbook.save(buffer)
        return SimpleUploadedFile(
            "carreras.xlsx", buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_only_owner_can_query_task_status(self):
        from rest_framework.test import APIClient

        self.client.force_authenticate(self.owner)
        accepted = self.client.post("/api/v1/import-export/import/carreras/", {"file": self._upload()}, format="multipart")
        self.assertEqual(accepted.status_code, 202, accepted.content)
        self.assertEqual(accepted.data["estado"], "pendiente")
        url = f"/api/v1/import-export/tareas/{accepted.data['task_id']}/"

        mine = self.client.get(url)
        self.assertEqual(mine.status_code, 200)
        self.assertIn(mine.data["estado"], {"completada", "fallida"})
        self.assertIsNotNone(mine.data["http_status"])

        other_client = APIClient()
        other_client.force_authenticate(self.other)
        self.assertEqual(other_client.get(url).status_code, 404)
        self.assertEqual(self.client.get("/api/v1/import-export/tareas/inexistente/").status_code, 404)

class CareerCatalogSpreadsheetTests(APITestCase):
    def test_import_preserves_manual_information_and_export_keeps_four_columns(self):
        province = Provincia.objects.create(nombre="Provincia catálogo")
        ces = Ces.objects.create(nombre="CES catálogo")
        career = Carrera.objects.create(
            codigo="CAT-1",
            nombre="Carrera catálogo",
            descripcion="Descripción inicial",
            campo_laboral="Campo inicial",
            requisitos="Requisitos iniciales",
            perfil_egresado="Perfil inicial",
            ces=ces,
            provincia=province,
        )
        service = CareerExcelService()

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Código", "Nombre", "CES", "Provincia"])
        worksheet.append([career.codigo, "Carrera actualizada", ces.nombre, province.nombre])
        upload = BytesIO()
        workbook.save(upload)
        upload.seek(0)

        result = service.import_file(upload)

        self.assertEqual(result.errors, [])
        career.refresh_from_db()
        self.assertEqual(career.nombre, "Carrera actualizada")
        self.assertEqual(career.descripcion, "Descripción inicial")
        self.assertEqual(career.campo_laboral, "Campo inicial")
        self.assertEqual(career.requisitos, "Requisitos iniciales")
        self.assertEqual(career.perfil_egresado, "Perfil inicial")

        exported = load_workbook(BytesIO(service.export_file(Carrera.objects.filter(pk=career.pk))))
        exported_sheet = exported.active
        self.assertEqual([cell.value for cell in exported_sheet[1]], ["Código", "Nombre", "CES", "Provincia"])
        self.assertEqual(exported_sheet.max_column, 4)
