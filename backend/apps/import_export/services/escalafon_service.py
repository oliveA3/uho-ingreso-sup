import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import Workbook, load_workbook

from apps.authentication.models import Estudiante
from apps.gestion_escuela.models import Escalafon, EscalafonItem
from apps.superadmin.models import Escuela


REQUIRED_HEADERS = [
    "CI", "Nombre", "Apellidos", "Sexo", "Dirección",
    "Índice_10mo", "Índice_11mo", "Índice_12mo", "Índice_General",
]


class EscalafonImportResult:
    def __init__(self, inserted=0, errors=None):
        self.inserted = inserted
        self.errors = errors or []


class EscalafonExcelService:
    def import_file(self, file_obj, escuela, proceso):
        workbook = load_workbook(filename=io.BytesIO(file_obj.read()), data_only=True)
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        missing = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing:
            return EscalafonImportResult(errors=[{"row": 1, "errors": [f"Faltan columnas: {', '.join(missing)}"]}])

        header_indexes = {header: headers.index(header) for header in REQUIRED_HEADERS}
        rows = []
        errors = []
        cis = set()
        for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in values):
                continue
            data = {header: values[index] for header, index in header_indexes.items()}
            row_errors = []
            ci = str(data["CI"]).strip() if data["CI"] is not None else ""
            if not ci.isdigit() or len(ci) != 11:
                row_errors.append("CI debe contener exactamente 11 dígitos numéricos.")
            elif ci in cis:
                row_errors.append("CI repetido dentro del Excel.")
            else:
                cis.add(ci)

            nombre = str(data["Nombre"]).strip() if data["Nombre"] is not None else ""
            apellidos = str(data["Apellidos"]).strip() if data["Apellidos"] is not None else ""
            direccion = str(data["Dirección"]).strip() if data["Dirección"] is not None else ""
            if not nombre:
                row_errors.append("Nombre es obligatorio.")
            if not apellidos:
                row_errors.append("Apellidos es obligatorio.")
            if not direccion:
                row_errors.append("Dirección es obligatoria.")

            sexo = str(data["Sexo"]).strip().upper() if data["Sexo"] is not None else ""
            if sexo not in {"M", "F"}:
                row_errors.append("Sexo debe ser M o F.")

            indices = {}
            for header, field in {
                "Índice_10mo": "indice_10",
                "Índice_11mo": "indice_11",
                "Índice_12mo": "indice_12",
                "Índice_General": "indice_general",
            }.items():
                value = data[header]
                try:
                    decimal_value = Decimal(str(value))
                    if decimal_value < 0 or decimal_value > 100:
                        raise ValueError
                    if decimal_value.as_tuple().exponent < -2:
                        raise ValueError
                    indices[field] = decimal_value.quantize(Decimal("0.01"))
                except (InvalidOperation, TypeError, ValueError):
                    row_errors.append(f"{header} debe estar entre 0.00 y 100.00, con máximo 2 decimales.")

            if row_errors:
                errors.append({"row": row_number, "errors": row_errors, "data": data})
            else:
                rows.append({"ci": ci, "nombre": nombre, "apellidos": apellidos, "sexo": sexo, "direccion": direccion, **indices})

        if errors:
            return EscalafonImportResult(errors=errors)

        with transaction.atomic():
            escalafon = Escalafon.objects.create(escuela=escuela, proceso=proceso)
            for row in rows:
                estudiante, created = Estudiante.objects.get_or_create(
                    ci=row["ci"],
                    defaults={
                        "nombre": row["nombre"], "apellidos": row["apellidos"], "sexo": row["sexo"],
                        "direccion": row["direccion"], "escuela": escuela,
                    },
                )
                if not created and estudiante.escuela_id != escuela.id:
                    raise ValueError(f"El CI {row['ci']} ya pertenece a otra escuela.")
                if not created:
                    for field in ("nombre", "apellidos", "sexo", "direccion"):
                        setattr(estudiante, field, row[field])
                    estudiante.save(update_fields=["nombre", "apellidos", "sexo", "direccion"])
                EscalafonItem.objects.create(
                    escalafon=escalafon,
                    estudiante=estudiante,
                    indice_10=row["indice_10"],
                    indice_11=row["indice_11"],
                    indice_12=row["indice_12"],
                    indice_general=row["indice_general"],
                )
        return EscalafonImportResult(inserted=len(rows))

    def export_file(self, queryset):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(REQUIRED_HEADERS)
        for entry in queryset.select_related("estudiante"):
            student = entry.estudiante
            worksheet.append([
                student.ci, student.nombre, student.apellidos, student.sexo, student.direccion,
                entry.indice_10, entry.indice_11, entry.indice_12, entry.indice_general,
            ])
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output.read()


def resolve_school(value):
    try:
        return Escuela.objects.get(pk=int(value))
    except (Escuela.DoesNotExist, TypeError, ValueError):
        return Escuela.objects.get(nombre__iexact=str(value).strip())
