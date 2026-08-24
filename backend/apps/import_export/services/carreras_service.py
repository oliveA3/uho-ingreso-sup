import io

from django.db import transaction
from openpyxl import Workbook, load_workbook

from apps.superadmin.models import Carrera, Ces, Provincia


REQUIRED_HEADERS = ["Código", "Nombre", "CES", "Provincia"]


class CareerImportResult:
    def __init__(self, inserted=0, errors=None):
        self.inserted = inserted
        self.errors = errors or []


class CareerExcelService:
    def import_file(self, file_obj):
        workbook = load_workbook(filename=io.BytesIO(file_obj.read()), data_only=True)
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        missing = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing:
            return CareerImportResult(errors=[{"row": 1, "errors": [f"Faltan columnas: {', '.join(missing)}"]}])

        rows = []
        errors = []
        codes = set()
        header_indexes = {header: headers.index(header) for header in REQUIRED_HEADERS}
        for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in values):
                continue
            data = {header: values[index] for header, index in header_indexes.items()}
            row_errors = []
            code = str(data["Código"]).strip() if data["Código"] is not None else ""
            name = str(data["Nombre"]).strip() if data["Nombre"] is not None else ""
            ces_name = str(data["CES"]).strip() if data["CES"] is not None else ""
            province_name = str(data["Provincia"]).strip() if data["Provincia"] is not None else ""
            if not code:
                row_errors.append("Código es obligatorio.")
            elif code in codes:
                row_errors.append("Código repetido dentro del Excel.")
            else:
                codes.add(code)
            if not name:
                row_errors.append("Nombre es obligatorio.")
            try:
                ces = Ces.objects.get(nombre__iexact=ces_name)
            except Ces.DoesNotExist:
                ces = None
                row_errors.append(f"CES no encontrado: {ces_name}.")
            try:
                province = Provincia.objects.get(nombre__iexact=province_name)
            except Provincia.DoesNotExist:
                province = None
                row_errors.append(f"Provincia no encontrada: {province_name}.")
            if row_errors:
                errors.append({"row": row_number, "errors": row_errors, "data": data})
            else:
                rows.append(Carrera(codigo=code, nombre=name, ces=ces, provincia=province, activa=True))

        if errors:
            return CareerImportResult(errors=errors)
        with transaction.atomic():
            Carrera.objects.all().delete()
            Carrera.objects.bulk_create(rows)
        return CareerImportResult(inserted=len(rows))

    def export_file(self, queryset):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(REQUIRED_HEADERS)
        for career in queryset.select_related("ces", "provincia"):
            worksheet.append([career.codigo, career.nombre, career.ces.nombre, career.provincia.nombre])
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output.read()
