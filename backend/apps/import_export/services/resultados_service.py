import io
import unicodedata
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook

from apps.authentication.models import Estudiante
from apps.gestion_personal.models import ConfirmacionPrueba, ResultadoExamen
from apps.gestion_provincial.models import ETAPAS_NOMBRES
from apps.superadmin.models import Asignatura


REQUIRED_HEADERS = ["CI", "Asignatura", "Nota"]
SUBJECTS = {
    "matematica": "Matematica",
    "espanol": "Espanol",
    "historia": "Historia",
}
SUBJECT_CATALOG_ALIASES = {
    "matematica": {"matematica"},
    "espanol": {"espanol"},
    "historia": {"historia", "historia de cuba"},
}


def normalize_subject(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(char for char in text if not unicodedata.combining(char))


class ResultadosImportResult:
    def __init__(self, inserted=0, updated=0, errors=None, students=None):
        self.inserted = inserted
        self.updated = updated
        self.errors = errors or []
        self.students = students or []


class ResultadosExcelService:
    def import_file(self, file_obj, proceso, provincia, selected_subject=None, deadline=None):
        selected_subject_key = normalize_subject(selected_subject) if selected_subject else None
        if selected_subject_key not in SUBJECTS:
            return ResultadosImportResult(errors=[{
                "row": 1,
                "errors": ["Debe seleccionar una asignatura válida: Matematica, Espanol o Historia."],
            }])
        if deadline is None:
            return ResultadosImportResult(errors=[{
                "row": 1,
                "errors": ["Debe indicar la fecha límite de reclamaciones."],
            }])
        workbook = load_workbook(filename=io.BytesIO(file_obj.read()), data_only=True)
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        missing = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing:
            return ResultadosImportResult(errors=[{
                "row": 1,
                "errors": [f"Faltan columnas: {', '.join(missing)}"],
            }])

        header_indexes = {header: headers.index(header) for header in REQUIRED_HEADERS}
        rows = []
        errors = []
        seen = set()
        students = {}
        for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in values):
                continue
            data = {header: values[index] for header, index in header_indexes.items()}
            row_errors = []
            ci = str(data["CI"]).strip() if data["CI"] is not None else ""
            subject_key = normalize_subject(data["Asignatura"])
            subject = SUBJECTS.get(subject_key)
            key = (ci, subject_key)
            if not ci.isdigit() or len(ci) != 11:
                row_errors.append("CI debe contener exactamente 11 dígitos numéricos.")
            elif key in seen:
                row_errors.append("La combinación CI y Asignatura está repetida dentro del Excel.")
            else:
                seen.add(key)
            if not subject:
                row_errors.append("Asignatura debe ser Matematica, Espanol o Historia.")
            elif subject_key != selected_subject_key:
                row_errors.append("La asignatura del Excel no coincide con la asignatura seleccionada.")

            try:
                grade = Decimal(str(data["Nota"]))
                if grade < 0 or grade > 100 or grade.as_tuple().exponent < -2:
                    raise ValueError
                grade = float(grade)
            except (InvalidOperation, TypeError, ValueError):
                row_errors.append("Nota debe estar entre 0.00 y 100.00, con máximo 2 decimales.")
                grade = None

            student_query = Estudiante.objects.filter(
                ci=ci,
                confirmacionprueba__proceso__anio__year=proceso.anio.year,
                confirmacionprueba__proceso__etapa__nombre=ETAPAS_NOMBRES[4],
            )
            if provincia is not None:
                student_query = student_query.filter(escuela__municipio__provincia=provincia)
            student = student_query.distinct().first() if ci.isdigit() and len(ci) == 11 else None
            if student is None:
                row_errors.append("El CI no existe en la base de datos del proceso y la provincia indicada.")

            if row_errors:
                errors.append({"row": row_number, "errors": row_errors, "data": data})
            else:
                rows.append((student, subject, grade))
                students[student.id] = student

        if errors:
            return ResultadosImportResult(errors=errors)

        subjects = {}
        for subject_key, name in SUBJECTS.items():
            subjects[name] = next(
                (
                    item for item in Asignatura.objects.filter(activa=True)
                    if normalize_subject(item.nombre) in SUBJECT_CATALOG_ALIASES[subject_key]
                ),
                None,
            )
        for index, (student, subject, grade) in enumerate(rows):
            if subjects[subject] is None:
                return ResultadosImportResult(errors=[{
                    "row": index + 2,
                    "errors": [f"No existe la asignatura '{subject}' en el catálogo."],
                }])

        inserted = updated = 0
        with transaction.atomic():
            for student, subject, grade in rows:
                subject_obj = subjects[subject]
                result = ResultadoExamen.objects.filter(
                    estudiante=student,
                    proceso=proceso,
                    asignatura=subject_obj,
                ).first()
                if result is None:
                    ResultadoExamen.objects.create(
                        estudiante=student,
                        proceso=proceso,
                        asignatura=subject_obj,
                        nota=grade,
                        fecha_limite_reclamo=deadline,
                    )
                    created = True
                else:
                    result.nota = grade
                    result.save(update_fields=["nota"])
                    created = False
                if created:
                    inserted += 1
                else:
                    updated += 1

        return ResultadosImportResult(inserted=inserted, updated=updated, students=list(students.values()))