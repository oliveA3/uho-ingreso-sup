from pathlib import Path

from django.conf import settings
from openpyxl import load_workbook


MAX_EXCEL_UPLOAD_BYTES = int(getattr(settings, "MAX_EXCEL_UPLOAD_BYTES", 10 * 1024 * 1024))
ALLOWED_EXCEL_EXTENSIONS = {".xlsx"}
ALLOWED_EXCEL_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def validate_excel_upload(upload):
    if not upload:
        return "Debe adjuntar un archivo Excel."
    max_mb = MAX_EXCEL_UPLOAD_BYTES // (1024 * 1024)
    if upload.size > MAX_EXCEL_UPLOAD_BYTES:
        return f"El archivo no puede superar los {max_mb} MB."
    if Path(upload.name).suffix.lower() not in ALLOWED_EXCEL_EXTENSIONS:
        return "Solo se permiten archivos Excel .xlsx."
    if (upload.content_type or "").lower() not in ALLOWED_EXCEL_MIME_TYPES:
        return "El tipo MIME del archivo no es un Excel permitido."
    try:
        upload.seek(0)
        workbook = load_workbook(filename=upload, read_only=True, data_only=True, keep_vba=False)
        if not workbook.sheetnames:
            return "El archivo Excel no contiene hojas."
        MAX_ROWS = 50000
        sheet = workbook.active
        if sheet.max_row and sheet.max_row > MAX_ROWS:
            return f"El archivo Excel no puede tener más de {MAX_ROWS} filas."
        workbook.close()
        upload.seek(0)
    except Exception:
        upload.seek(0)
        return "El archivo no es un Excel válido o está dañado."
    return None
