"""Servicios de importación Excel ejecutados de forma asíncrona (Celery).

Las vistas validan permisos/parámetros, guardan el archivo en el storage y encolan
`run_import`; la tarea lee el Excel, valida y escribe de forma atómica (todo o nada).
El resultado es `{"http_status": int, "body": {...}}` (mismo cuerpo que antes devolvía la vista).
"""
import io
import logging
import unicodedata
from datetime import date
from decimal import Decimal, InvalidOperation

from celery import shared_task
from django.core.files.storage import default_storage
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.authentication.models import Estudiante, Usuario
from apps.core.audit import record_audit
from apps.gestion_provincial.models import ETAPAS_NOMBRES, CorteCarrera, Etapa, Otorgamiento, PlanPlaza, Proceso
from apps.superadmin.models import Carrera, Ces, Escuela, Provincia, TipoOtorgamiento
from .carreras_service import CareerExcelService
from .escalafon_service import EscalafonExcelService
from .resultados_service import ResultadosExcelService

logger = logging.getLogger("django.request")

IMPORT_KINDS = ("plan_plaza", "otorgamiento", "corte", "carreras", "escalafon", "resultados")
GENERIC_ERROR = {"detail": "No se pudo procesar el archivo Excel. Verifique que el formato sea válido e inténtelo de nuevo."}


class _AuditRequest:
    """Mínimo que necesita `record_audit` (IP) fuera de una petición HTTP."""

    def __init__(self, path, ip):
        self.path = path
        self.META = {"REMOTE_ADDR": ip or ""}


def _result(status, body):
    return {"http_status": status, "body": body}


def _normalize_header(value):
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text.replace("_", " ").replace("-", " ").replace(".", " ").replace("/", " ")


def _get_active_process():
    stage = Etapa.objects.filter(nombre=ETAPAS_NOMBRES[3]).first()
    if not stage:
        raise ValueError("No existe la etapa para planes de plaza.")
    process = Proceso.objects.filter(anio__year=timezone.now().year, etapa=stage).order_by("-id").first()
    if process is None:
        process = Proceso.objects.create(
            anio=timezone.now().date().replace(month=1, day=1),
            etapa=stage,
        )
    return process


def import_plan_plaza(file, user, audit_request):
    required = ["Codigo_Carrera", "Nombre_Carrera", "Cantidad_Plazas", "Tipo_Otorgamiento", "CES", "Provincia", "Sexo"]
    try:
        workbook = load_workbook(file, data_only=True)
        sheet = workbook.active
        raw_headers = [cell.value for cell in sheet[1]]
    except Exception:
        logger.exception("No se pudo leer el archivo Excel de plan de plazas.")
        return _result(400, {"detail": "No se pudo leer el archivo Excel. Verifique que el formato sea válido."})

    normalized_headers = {_normalize_header(header): header for header in raw_headers if header is not None}
    missing = [header for header in required if _normalize_header(header) not in normalized_headers]
    if missing:
        return _result(400, {"detail": "Faltan columnas requeridas.", "missing": missing})

    process = _get_active_process()
    result = {"inserted": 0, "updated": 0, "errors": []}
    validated_rows = []

    user_rol = getattr(user, "rol", None)
    user_scoped_roles = {"jefe_comision", "ingreso_provincial"}

    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        data = {}
        for index, header in enumerate(raw_headers):
            if index < len(values):
                key = _normalize_header(header)
                data[key] = values[index]
        try:
            codigo = str(data.get("codigo carrera") or "").strip()
            nombre = str(data.get("nombre carrera") or "").strip()
            carrera = Carrera.objects.filter(codigo__iexact=codigo).first()
            if carrera is None and nombre:
                carrera = Carrera.objects.filter(nombre__iexact=nombre).first()
            if carrera is None:
                raise ValueError(f"No existe una carrera con código o nombre '{codigo or nombre}'.")

            ces_name = str(data.get("ces") or "").strip()
            ces = Ces.objects.filter(nombre__iexact=ces_name).first() if ces_name else None
            if ces is None and carrera.ces_id:
                ces = carrera.ces
            if ces is None:
                raise ValueError(f"No existe el CES '{ces_name or 'vacío'}'.")

            row_provincia = str(data.get("provincia") or "").strip()
            universidad_provincia = Provincia.objects.filter(nombre__iexact=row_provincia).first() if row_provincia else None
            if universidad_provincia is None and carrera.provincia_id:
                universidad_provincia = carrera.provincia

            if user_rol in user_scoped_roles and user.provincia_id:
                provincia = user.provincia
            else:
                provincia = universidad_provincia

            if provincia is None:
                raise ValueError(f"No existe la provincia '{row_provincia or 'asociada al usuario'}'.")

            amount = int(data.get("cantidad plazas"))
            if amount <= 0:
                raise ValueError("Cantidad_Plazas debe ser un entero positivo.")

            tipo_name = str(data.get("tipo otorgamiento") or "").strip()
            tipo = TipoOtorgamiento.objects.filter(nombre__iexact=tipo_name, activa=True).first()
            if tipo is None:
                raise ValueError("Tipo_Otorgamiento debe ser un tipo de otorgamiento activo (p. ej. 'Municipal' o 'Provincial').")

            sex = str(data.get("sexo") or "").strip().upper()
            if sex not in {"A", "F", "M"}:
                raise ValueError("Sexo debe ser 'A', 'F' o 'M'.")

            if nombre and str(carrera.nombre).strip().lower() != str(nombre).strip().lower():
                raise ValueError("Nombre_Carrera no coincide con la carrera encontrada.")
            if carrera.ces_id != ces.id:
                raise ValueError("El CES no coincide con la carrera.")

            validated_rows.append((carrera, sex, amount, tipo, ces, provincia))
        except ValueError as error:
            result["errors"].append({
                "row": row_number,
                "error": str(error),
                "data": {key: value for key, value in (data or {}).items() if key is not None},
            })
        except Exception:
            logger.exception("Error inesperado procesando la fila %s de importación de plan de plazas.", row_number)
            result["errors"].append({
                "row": row_number,
                "error": "Error inesperado al procesar esta fila.",
                "data": {key: value for key, value in (data or {}).items() if key is not None},
            })

    if result["errors"]:
        return _result(400, result)

    with transaction.atomic():
        for carrera, sex, amount, tipo, ces, provincia in validated_rows:
            _, created = PlanPlaza.objects.update_or_create(
                proceso=process,
                carrera=carrera,
                sexo=sex,
                defaults={
                    "cantidad_plazas": amount,
                    "otorgamiento_tipo": tipo,
                    "ces": ces,
                    "provincia": provincia,
                },
            )
            result["inserted" if created else "updated"] += 1

    record_audit(user, "Importación de plan de plazas", audit_request.path, request=audit_request, new={"inserted": result["inserted"], "updated": result["updated"], "proceso": process.id})
    return _result(201, result)


def _normalize_import_header(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return "".join(char for char in text if char.isalnum())


def _get_stage_six_process():
    stage = Etapa.objects.filter(nombre=ETAPAS_NOMBRES[6]).first()
    process = Proceso.get_for_stage_and_year(timezone.now().year, ETAPAS_NOMBRES[6])
    if not stage or not process:
        raise ValueError("No existe un proceso de otorgamiento para el año actual.")
    if stage.estado != "en_curso":
        raise ValueError("La importación solo está disponible durante la etapa 6.")
    return process


def _read_stage_six_workbook(file_obj, required_headers):
    workbook = load_workbook(filename=file_obj, data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    normalized = {_normalize_import_header(header): header for header in headers if header is not None}
    missing = [header for header in required_headers if _normalize_import_header(header) not in normalized]
    if missing:
        raise ValueError(f"Faltan columnas requeridas: {', '.join(missing)}")
    return worksheet, headers, normalized


def _resolve_stage_six_career(data):
    codigo = str(data.get("codigo_carrera") or "").strip()
    nombre = str(data.get("nombre_carrera") or "").strip()
    carrera = Carrera.objects.filter(codigo__iexact=codigo).first()
    if not carrera:
        raise ValueError(f"No existe una carrera con código '{codigo}'.")
    if not nombre or carrera.nombre.strip().lower() != nombre.lower():
        raise ValueError("Nombre_Carrera no coincide con la carrera del código indicado.")
    return carrera


def import_stage_six(file_obj, kind):
    required_headers = (
        ["CI", "Codigo_Carrera", "Nombre_Carrera", "Indice_Otorgamiento"]
        if kind == "otorgamiento"
        else ["Codigo_Carrera", "Nombre_Carrera", "Indice_Corte"]
    )
    try:
        process = _get_stage_six_process()
        worksheet, headers, normalized_headers = _read_stage_six_workbook(file_obj, required_headers)
    except ValueError as error:
        return {"inserted": 0, "updated": 0, "errors": [{"row": 1, "errors": [str(error)]}]}
    except Exception:
        logger.exception("No se pudo leer el archivo Excel de etapa 6 (%s).", kind)
        return {"inserted": 0, "updated": 0, "errors": [{"row": 1, "errors": ["No se pudo leer el archivo Excel. Verifique que el formato sea válido."]}]}

    rows = []
    errors = []
    seen = set()
    for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        raw = {normalized_headers[_normalize_import_header(header)]: values[index] for index, header in enumerate(headers) if header is not None and index < len(values)}
        data = {
            "ci": raw.get(normalized_headers.get("ci")),
            "codigo_carrera": raw.get(normalized_headers.get("codigocarrera")),
            "nombre_carrera": raw.get(normalized_headers.get("nombrecarrera")),
            "indice": raw.get(normalized_headers.get("indiceotorgamiento" if kind == "otorgamiento" else "indicecorte")),
        }
        row_errors = []
        student = None
        carrera = None
        if kind == "otorgamiento":
            ci = str(data["ci"] or "").strip()
            if not ci.isdigit() or len(ci) != 11:
                row_errors.append("CI debe contener exactamente 11 dígitos numéricos.")
            else:
                student = Estudiante.objects.filter(ci=ci).first()
                if not student:
                    row_errors.append("El CI no existe en la base de datos.")
        try:
            carrera = _resolve_stage_six_career(data)
        except ValueError as error:
            row_errors.append(str(error))
        try:
            indice = Decimal(str(data["indice"]))
            if indice < 0 or indice > 100 or indice.as_tuple().exponent < -2:
                raise ValueError
            data["indice"] = float(indice)
        except (InvalidOperation, TypeError, ValueError):
            row_errors.append("El índice debe ser un decimal entre 0 y 100, con máximo 2 decimales.")
        key = (student.id if student else data["ci"], carrera.id if carrera else data["codigo_carrera"])
        if key in seen:
            row_errors.append("La combinación indicada está repetida dentro del Excel.")
        else:
            seen.add(key)
        if row_errors:
            errors.append({"row": row_number, "errors": row_errors, "data": raw})
        else:
            rows.append((student, carrera, data["indice"]))
    if errors:
        return {"inserted": 0, "updated": 0, "errors": errors}

    inserted = updated = 0
    with transaction.atomic():
        for student, carrera, indice in rows:
            lookup = {"proceso": process, "carrera": carrera}
            if kind == "otorgamiento":
                lookup["estudiante"] = student
                _, created = Otorgamiento.objects.update_or_create(**lookup, defaults={"indice_otorgamiento": indice})
            else:
                _, created = CorteCarrera.objects.update_or_create(**lookup, defaults={"indice_corte": indice})
            if created:
                inserted += 1
            else:
                updated += 1
    return {"inserted": inserted, "updated": updated, "errors": []}


def import_otorgamiento(file, user, audit_request):
    result = import_stage_six(file, "otorgamiento")
    if not result["errors"]:
        record_audit(user, "Importación de otorgamientos", audit_request.path, request=audit_request, new={"inserted": result["inserted"], "updated": result["updated"]})
        from apps.core.notifications import notify_users
        students = Usuario.objects.filter(
            rol="estudiante",
            escuela__municipio__provincia=user.provincia,
        ) if user.rol == "jefe_comision" else Usuario.objects.filter(rol="estudiante")
        notify_users(
            students,
            "Otorgamientos publicados",
            f"Se publicaron los otorgamientos de carreras del proceso {timezone.now().year}. Ya puedes consultar tu resultado.",
            filter_students=False,
        )
    return _result(400 if result["errors"] else 201, result)


def import_corte(file, user, audit_request):
    result = import_stage_six(file, "corte")
    if not result["errors"]:
        record_audit(user, "Importación de índices de corte", audit_request.path, request=audit_request, new={"inserted": result["inserted"], "updated": result["updated"]})
    return _result(400 if result["errors"] else 201, result)


def import_carreras(file, user, audit_request):
    result = CareerExcelService().import_file(file)
    if result.errors:
        return _result(400, {"inserted": 0, "errors": result.errors})
    return _result(200, {"inserted": result.inserted, "errors": []})


def import_escalafon(file, user, audit_request, params):
    escuela = Escuela.objects.select_related("municipio").get(pk=params["escuela_id"])
    proceso = Proceso.objects.get(pk=params["proceso_id"])
    anio = params["anio"]
    try:
        result = EscalafonExcelService().import_file(file, escuela, proceso)
    except (IntegrityError, ValueError) as error:
        if isinstance(error, IntegrityError):
            error = "Ya existe un escalafón para esa escuela y año."
        return _result(400, {"detail": str(error)})
    if result.errors:
        return _result(400, {"inserted": 0, "errors": result.errors})
    from apps.core.notifications import notify_users
    notify_users(
        Usuario.objects.filter(rol="estudiante", escuela=escuela),
        "Escalafón actualizado",
        f"El secretario publicó el escalafón de tu escuela para el proceso {anio}.",
    )
    return _result(201, {"inserted": result.inserted, "errors": []})


def import_resultados(file, user, audit_request, params):
    proceso = Proceso.objects.select_related("etapa").get(pk=params["proceso_id"])
    selected_subject = params.get("asignatura")
    deadline = date.fromisoformat(params["deadline"])
    try:
        result = ResultadosExcelService().import_file(
            file,
            proceso,
            user.provincia,
            selected_subject=selected_subject,
            deadline=deadline,
        )
    except ValueError as error:
        return _result(400, {"detail": str(error)})
    except Exception:
        logger.exception("No se pudo procesar el archivo Excel de resultados.")
        return _result(400, {"detail": "No se pudo leer el archivo Excel. Verifique que el formato sea válido."})
    if result.errors:
        return _result(400, {"inserted": 0, "updated": 0, "errors": result.errors})
    from apps.core.notifications import notify_users
    notification_users = Usuario.objects.filter(rol="estudiante")
    if user.rol != "superadmin":
        notification_users = notification_users.filter(escuela__municipio__provincia=user.provincia)
    notify_users(
        notification_users,
        "Resultados de exámenes publicados",
        f"Ya puedes consultar tus resultados de {selected_subject} del proceso {proceso.anio.year}.",
        filter_students=False,
    )
    record_audit(user, "Importación de resultados de exámenes", audit_request.path, request=audit_request, new={"inserted": result.inserted, "updated": result.updated, "anio": proceso.anio.year, "asignatura": selected_subject})
    return _result(201, {"inserted": result.inserted, "updated": result.updated, "errors": []})


@shared_task(name="import_export.run_import")
def run_import(kind, storage_path, params, user_id):
    """Ejecuta una importación Excel. Nunca propaga excepciones: devuelve `http_status` y `body`."""
    params = params or {}
    try:
        user = Usuario.objects.get(pk=user_id)
        audit_request = _AuditRequest(params.get("path", ""), params.get("ip", ""))
        with default_storage.open(storage_path, "rb") as stored:
            file = io.BytesIO(stored.read())
        if kind == "plan_plaza":
            return import_plan_plaza(file, user, audit_request)
        if kind == "otorgamiento":
            return import_otorgamiento(file, user, audit_request)
        if kind == "corte":
            return import_corte(file, user, audit_request)
        if kind == "carreras":
            return import_carreras(file, user, audit_request)
        if kind == "escalafon":
            return import_escalafon(file, user, audit_request, params)
        if kind == "resultados":
            return import_resultados(file, user, audit_request, params)
        return _result(400, {"detail": "Tipo de importación no soportado."})
    except Exception:
        logger.exception("Error inesperado en la importación asíncrona (%s).", kind)
        return _result(500, dict(GENERIC_ERROR))
    finally:
        try:
            default_storage.delete(storage_path)
        except Exception:
            logger.warning("No se pudo eliminar el archivo temporal de importación %s.", storage_path)


def enqueue_import(request, kind, file, params=None):
    """Guarda el archivo, registra el dueño de la tarea y encola `run_import`. Devuelve el task_id."""
    import uuid
    from django.conf import settings
    from django.core.cache import cache
    from django.core.files.base import ContentFile

    file.seek(0)
    storage_path = default_storage.save(f"imports/{uuid.uuid4().hex}.xlsx", ContentFile(file.read()))
    task_id = str(uuid.uuid4())
    cache.set(f"import_task:{task_id}", {"user_id": request.user.id, "kind": kind}, 3600)
    params = dict(params or {}, path=request.path, ip=request.META.get("REMOTE_ADDR", ""))
    async_result = run_import.apply_async(args=[kind, storage_path, params, request.user.id], task_id=task_id)
    if getattr(settings, "CELERY_TASK_ALWAYS_EAGER", False):
        # En modo eager no hay backend de resultados: se deja el resultado en caché.
        cache.set(f"import_result:{task_id}", async_result.result, 3600)
    return task_id
