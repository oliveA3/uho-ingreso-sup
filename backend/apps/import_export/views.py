import logging
import unicodedata

from rest_framework.views import APIView
from rest_framework.response import Response
from django.http import HttpResponse
from django.db.models import Count, Q
from drf_spectacular.utils import extend_schema, inline_serializer, OpenApiParameter, OpenApiExample
from drf_spectacular.types import OpenApiTypes
from rest_framework import status, permissions, serializers
from django.db import transaction
from apps.gestion_provincial.models import PlanPlaza, Otorgamiento, CorteCarrera, Proceso, Etapa
from apps.authentication.models import Usuario
from apps.core.audit import record_audit
from apps.core.throttling import BulkOperationRateThrottle
from apps.gestion_provincial.permissions import IsCareerManager
from apps.superadmin.models import Asignatura, Carrera, Ces, Escuela, Municipio, Provincia, TipoOtorgamiento
from .services.carreras_service import CareerExcelService
from .services.escalafon_service import EscalafonExcelService, rank_escalafon_entries, resolve_school
from apps.gestion_escuela.models import Escalafon, EscalafonItem
from apps.gestion_escuela.serializers import EscalafonItemSerializer, StudentEscalafonActionSerializer
from apps.gestion_escuela.permissions import CanManageEscalafon
from apps.gestion_provincial.models import ETAPAS_NOMBRES, Etapa
from django.utils import timezone
import io
from datetime import date, datetime

from openpyxl import Workbook
from apps.gestion_personal.models import BoletaInteres, BoletaSolicitud, BoletaSolicitudItem, Reclamacion
from apps.gestion_personal.models import ResultadoExamen
from .services.upload_validation import validate_excel_upload
from .services.import_tasks import enqueue_import
from apps.core.pdf import render_pdf

logger = logging.getLogger("django.request")


_IMPORT_ACCEPTED = inline_serializer(
    name="ImportTaskAcceptedResponse",
    fields={
        "task_id": serializers.CharField(help_text="Identificador de la tarea; consultar en /import-export/tareas/{task_id}/."),
        "estado": serializers.CharField(help_text="Siempre 'pendiente' al encolar."),
    },
)


def _accepted(task_id):
    return Response({"task_id": task_id, "estado": "pendiente"}, status=status.HTTP_202_ACCEPTED)


def build_pdf_document(title, lines, process_name=None, process_identifier=None):
    header_lines = []
    if process_name:
        header_lines.append(f"Proceso: {process_name}")
    if process_identifier:
        header_lines.append(f"Identificación del proceso: {process_identifier}")
    return render_pdf(title, lines, header_lines)


def _ballot_pdf(title, student, items, career_getter, process_name=None, process_identifier=None):
    escalafon_entry = EscalafonItem.objects.filter(
        estudiante=student,
        escalafon__proceso__anio__year=timezone.now().year,
    ).order_by("-escalafon__proceso__anio", "-escalafon_id", "-id").first()
    index = escalafon_entry.indice_general if escalafon_entry else student.indice_general
    lines = [
        f"Nombre: {student.nombre} {student.apellidos}",
        f"CI: {student.ci}",
        f"Escuela: {student.escuela.nombre}",
        f"Índice general: {index or ''}",
        "",
        "Prioridad | Carrera | CES | Provincia universidad",
    ]
    for item in items:
        career = career_getter(item)
        lines.append(f"{item.prioridad} | {career.nombre} | {career.ces.nombre} | {career.provincia.nombre}")
    return build_pdf_document(title, lines, process_name=process_name, process_identifier=process_identifier)


def _ballot_excel(title, student, items, career_getter, process):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Boleta"
    sheet.append([title])
    sheet.append(["Nombre", f"{student.nombre} {student.apellidos}"])
    sheet.append(["CI", student.ci])
    sheet.append(["Escuela", student.escuela.nombre])
    process_label = process.etapa.nombre if process and process.etapa else "Proceso de ingreso"
    sheet.append(["Proceso", process_label])
    sheet.append([])
    sheet.append(["Prioridad", "Código", "Carrera", "CES", "Provincia universidad"])
    for item in items:
        career = career_getter(item)
        sheet.append([
            item.prioridad,
            career.codigo,
            career.nombre,
            career.ces.nombre,
            career.provincia.nombre,
        ])

    sheet.freeze_panes = "A7"
    sheet.column_dimensions["A"].width = 12
    for column in ("B", "C", "D", "E"):
        sheet.column_dimensions[column].width = 28
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@extend_schema(
    tags=["Importación y exportación"],
    summary="Descargar boleta de interés en PDF",
    description=(
        "Genera y descarga en PDF la boleta de interés del estudiante autenticado para el proceso de ingreso "
        "vigente (etapa 'Boleta de interés'). Solo accesible para usuarios con rol 'estudiante'."
    ),
    responses={(200, "application/pdf"): OpenApiTypes.BINARY},
)
class StudentInterestPdfExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if request.user.rol != "estudiante":
            return Response({"detail": "Solo un estudiante puede descargar esta boleta."}, status=403)
        student = getattr(request.user, "estudiante", None)
        process = Proceso.get_for_stage_and_year(timezone.now().year, ETAPAS_NOMBRES[2])
        ballot = BoletaInteres.objects.filter(estudiante=student, proceso=process).first() if student and process else None
        if not ballot:
            return Response({"detail": "No existe una boleta de interés."}, status=404)
        items = ballot.boleta_interes.select_related("carrera__ces", "carrera__provincia").order_by("prioridad")
        process_label = process.nombre if process else "Proceso de ingreso"
        response = HttpResponse(
            _ballot_pdf(
                "BOLETA DE INTERES",
                student,
                items,
                lambda item: item.carrera,
                process_name=f"Proceso de Ingreso {timezone.now().year}",
                process_identifier=f"{process_label}",
            ),
            content_type="application/pdf",
        )
        response["Content-Disposition"] = 'attachment; filename="boleta-interes.pdf"'
        return response


@extend_schema(
    tags=["Importación y exportación"],
    summary="Descargar boleta de interés en Excel",
    description=(
        "Genera y descarga en Excel (.xlsx) la boleta de interés del estudiante autenticado para el proceso de "
        "ingreso vigente. Solo accesible para usuarios con rol 'estudiante'."
    ),
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class StudentInterestExcelExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if request.user.rol != "estudiante":
            return Response({"detail": "Solo un estudiante puede descargar esta boleta."}, status=403)
        student = getattr(request.user, "estudiante", None)
        process = Proceso.get_for_stage_and_year(timezone.now().year, ETAPAS_NOMBRES[2])
        ballot = BoletaInteres.objects.filter(estudiante=student, proceso=process).first() if student and process else None
        if not ballot:
            return Response({"detail": "No existe una boleta de interés."}, status=404)
        items = ballot.boleta_interes.select_related("carrera__ces", "carrera__provincia").order_by("prioridad")
        response = HttpResponse(
            _ballot_excel("BOLETA DE INTERES", student, items, lambda item: item.carrera, process),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="boleta-interes.xlsx"'
        return response


@extend_schema(
    tags=["Importación y exportación"],
    summary="Descargar boleta de solicitud en PDF",
    description=(
        "Genera y descarga en PDF la boleta de solicitud. Un estudiante autenticado descarga su propia boleta "
        "del proceso vigente (sin indicar `ballot_id`); un Secretario o Director de escuela puede descargar la "
        "boleta de un estudiante de su escuela indicando `ballot_id` en la ruta."
    ),
    responses={(200, "application/pdf"): OpenApiTypes.BINARY},
)
class SolicitudPdfExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, ballot_id=None):
        student = getattr(request.user, "estudiante", None)
        if request.user.rol == "estudiante":
            process = Proceso.get_for_stage_and_year(timezone.now().year, ETAPAS_NOMBRES[3])
            ballot = BoletaSolicitud.objects.filter(estudiante=student, proceso=process).select_related("estudiante__escuela").first() if student and process else None
        elif request.user.rol in {"secretario_escuela", "director_escuela"} and request.user.escuela_id and ballot_id:
            ballot = BoletaSolicitud.objects.filter(pk=ballot_id, estudiante__escuela_id=request.user.escuela_id).select_related("estudiante__escuela").first()
        else:
            return Response({"detail": "No tienes permiso para descargar esta boleta."}, status=403)
        if not ballot:
            return Response({"detail": "No existe la boleta de solicitud."}, status=404)
        items = ballot.boleta_solicitud.select_related("plan_plaza__carrera__ces", "plan_plaza__carrera__provincia").order_by("prioridad")
        process_label = ballot.proceso.etapa.nombre if ballot.proceso and ballot.proceso.etapa else "Proceso de ingreso"
        response = HttpResponse(
            _ballot_pdf(
                "BOLETA DE SOLICITUD",
                ballot.estudiante,
                items,
                lambda item: item.plan_plaza.carrera,
                process_name=f"Proceso de Ingreso {timezone.now().year}",
                process_identifier=f"{process_label}",
            ),
            content_type="application/pdf",
        )
        response["Content-Disposition"] = 'attachment; filename="boleta-solicitud.pdf"'
        return response


@extend_schema(
    tags=["Importación y exportación"],
    summary="Descargar boleta de solicitud en Excel",
    description=(
        "Genera y descarga en Excel (.xlsx) la boleta de solicitud. Un estudiante autenticado descarga su propia "
        "boleta del proceso vigente (sin indicar `ballot_id`); un Secretario o Director de escuela puede descargar "
        "la boleta de un estudiante de su escuela indicando `ballot_id` en la ruta."
    ),
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class SolicitudExcelExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, ballot_id=None):
        student = getattr(request.user, "estudiante", None)
        if request.user.rol == "estudiante":
            process = Proceso.get_for_stage_and_year(timezone.now().year, ETAPAS_NOMBRES[3])
            ballot = BoletaSolicitud.objects.filter(estudiante=student, proceso=process).select_related("estudiante__escuela").first() if student and process else None
        elif request.user.rol in {"secretario_escuela", "director_escuela"} and request.user.escuela_id and ballot_id:
            ballot = BoletaSolicitud.objects.filter(pk=ballot_id, estudiante__escuela_id=request.user.escuela_id).select_related("estudiante__escuela").first()
        else:
            return Response({"detail": "No tienes permiso para descargar esta boleta."}, status=403)
        if not ballot:
            return Response({"detail": "No existe la boleta de solicitud."}, status=404)
        items = ballot.boleta_solicitud.select_related("plan_plaza__carrera__ces", "plan_plaza__carrera__provincia").order_by("prioridad")
        response = HttpResponse(
            _ballot_excel("BOLETA DE SOLICITUD", ballot.estudiante, items, lambda item: item.plan_plaza.carrera, ballot.proceso),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="boleta-solicitud.xlsx"'
        return response


def _disambiguate_ballot_operation_id(view_cls, bulk_operation_id, by_id_operation_id):
    """
    `SolicitudPdfExportView` and `SolicitudExcelExportView` are each registered on two
    routes in urls.py (one without `ballot_id` for the student's own ballot, one with
    `ballot_id` for a school staff member downloading a specific ballot). Since both
    routes point at the same view/method, drf-spectacular would otherwise generate the
    same operationId for both, producing a collision warning. This assigns distinct,
    explicit operationIds based on whether `ballot_id` is part of the resolved path.
    """
    base_schema_class = type(view_cls.schema)

    def get_operation_id(self):
        return by_id_operation_id if "ballot_id" in self.path else bulk_operation_id

    view_cls.schema = type(
        f"{view_cls.__name__}OperationIdSchema",
        (base_schema_class,),
        {"get_operation_id": get_operation_id},
    )()


_disambiguate_ballot_operation_id(
    SolicitudPdfExportView,
    bulk_operation_id="import_export_boleta_solicitud_pdf_bulk",
    by_id_operation_id="import_export_boleta_solicitud_pdf_by_id",
)
_disambiguate_ballot_operation_id(
    SolicitudExcelExportView,
    bulk_operation_id="import_export_boleta_solicitud_excel_bulk",
    by_id_operation_id="import_export_boleta_solicitud_excel_by_id",
)


@extend_schema(
    tags=["Importación y exportación"],
    summary="Importar plan de plazas desde Excel",
    description=(
        "Sube un archivo Excel (.xlsx) con el plan de plazas por carrera, tipo de otorgamiento, CES, provincia y "
        "sexo, y lo inserta o actualiza para el proceso de plan de plazas del año en curso. Columnas requeridas: "
        "Codigo_Carrera, Nombre_Carrera, Cantidad_Plazas, Tipo_Otorgamiento, CES, Provincia, Sexo. Requiere permiso "
        "de gestor de carreras (`IsCareerManager`). Los roles con alcance provincial ('jefe_comision', "
        "'ingreso_provincial') solo pueden registrar plazas en su propia provincia. La importación es atómica: si "
        "alguna fila tiene errores, no se guarda ningún registro.",
    ),
    request=inline_serializer(
        name="ImportPlanPlazaUploadRequest",
        fields={"file": serializers.FileField(help_text="Archivo .xlsx con el plan de plazas.")},
    ),
    responses={
        202: _IMPORT_ACCEPTED,
        201: inline_serializer(
            name="ImportPlanPlazaResponse",
            fields={
                "inserted": serializers.IntegerField(),
                "updated": serializers.IntegerField(),
                "errors": serializers.ListField(child=serializers.DictField(), default=list),
            },
        ),
        400: inline_serializer(
            name="ImportPlanPlazaErrorResponse",
            fields={
                "inserted": serializers.IntegerField(),
                "updated": serializers.IntegerField(),
                "errors": serializers.ListField(child=serializers.DictField()),
            },
        ),
    },
    examples=[
        OpenApiExample(
            "Importación exitosa",
            value={"success": True, "data": {"inserted": 12, "updated": 3, "errors": []}, "error": None},
            response_only=True,
        ),
    ],
)
class ImportPlanPlazaView(APIView):
    permission_classes = [IsCareerManager]
    throttle_classes = [BulkOperationRateThrottle]

    def post(self, request):
        file = request.FILES.get("file")
        upload_error = validate_excel_upload(file)
        if upload_error:
            return Response({"detail": upload_error}, status=400)
        return _accepted(enqueue_import(request, "plan_plaza", file))


def _normalize_plan_plaza_filename(provincia, anio):
    import re
    import unicodedata

    text = (provincia or "todos").strip()
    text = unicodedata.normalize("NFKD", text)
    text = ''.join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-zA-Z0-9]+", "_", text.lower()).strip("_") or "todos"
    return f"plan_plazas_{text}_{anio}.xlsx"


@extend_schema(
    tags=["Importación y exportación"],
    summary="Descargar plantilla de plan de plazas",
    description=(
        "Genera una plantilla Excel (.xlsx) de ejemplo con las columnas requeridas para importar el plan de "
        "plazas, prellenada con hasta 25 carreras activas de muestra. Requiere permiso de gestor de carreras."
    ),
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class PlanPlazaTemplateView(APIView):
    permission_classes = [IsCareerManager]

    def get(self, request):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plan de Plazas"
        sheet.append(["Codigo_Carrera", "Nombre_Carrera", "Cantidad_Plazas", "Tipo_Otorgamiento", "CES", "Provincia", "Sexo"])
        carreras = list(Carrera.objects.filter(activa=True).select_related("ces", "provincia").order_by("nombre")[:25])
        for index, carrera in enumerate(carreras, start=1):
            sheet.append([carrera.codigo, carrera.nombre, 10 + (index % 5), "Municipal" if index % 2 == 0 else "Provincial", carrera.ces.nombre, carrera.provincia.nombre, "A" if index % 3 else "F"])
        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="plan-plazas-prueba.xlsx"'
        return response


@extend_schema(
    tags=["Importación y exportación"],
    summary="Exportar plan de plazas",
    description="Exporta a Excel (.xlsx) el plan de plazas del proceso de plan de plazas, filtrable por año y provincia.",
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso (alias: year). Por defecto el año actual."),
        OpenApiParameter("provincia", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre exacto de la provincia para filtrar (alias: province)."),
    ],
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class ImportPlanPlazaExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        anio = request.query_params.get("anio") or request.query_params.get("year") or timezone.now().year
        provincia = request.query_params.get("provincia") or request.query_params.get("province") or ""

        queryset = PlanPlaza.objects.select_related("carrera", "carrera__ces", "carrera__provincia", "ces", "provincia", "proceso").filter(
            proceso__etapa__nombre=ETAPAS_NOMBRES[3],
            proceso__anio__year=anio,
        )
        if provincia:
            queryset = queryset.filter(provincia__nombre__iexact=provincia)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plan de Plazas"
        headers = [
            "Codigo_Carrera",
            "Nombre_Carrera",
            "Cantidad_Plazas",
            "Tipo_Otorgamiento",
            "CES",
            "Provincia",
            "Sexo",
        ]
        sheet.append(headers)

        for item in queryset.select_related("otorgamiento_tipo").order_by("carrera__nombre"):
            sheet.append([
                item.carrera.codigo,
                item.carrera.nombre,
                item.cantidad_plazas,
                item.otorgamiento_tipo.nombre,
                item.ces.nombre,
                item.carrera.provincia.nombre,
                item.sexo,
            ])

        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{_normalize_plan_plaza_filename(provincia or "todos", anio)}"'
        return response

@extend_schema(
    tags=["Importación y exportación"],
    summary="Importar otorgamientos de carreras",
    description=(
        "Sube un archivo Excel (.xlsx) con los otorgamientos de carreras (CI, Codigo_Carrera, Nombre_Carrera, "
        "Indice_Otorgamiento) durante la etapa 6 del proceso de ingreso en curso. Solo puede ejecutarlo el Jefe de "
        "Comisión (o superusuario). La importación es atómica y, al finalizar sin errores, notifica a los "
        "estudiantes afectados (de la provincia del jefe de comisión, o a todos si es superusuario).",
    ),
    request=inline_serializer(
        name="ImportOtorgamientoUploadRequest",
        fields={"file": serializers.FileField(help_text="Archivo .xlsx con los otorgamientos.")},
    ),
    responses={
        202: _IMPORT_ACCEPTED,
        201: inline_serializer(
            name="ImportOtorgamientoResponse",
            fields={
                "inserted": serializers.IntegerField(),
                "updated": serializers.IntegerField(),
                "errors": serializers.ListField(child=serializers.DictField(), default=list),
            },
        ),
        400: inline_serializer(
            name="ImportOtorgamientoErrorResponse",
            fields={
                "inserted": serializers.IntegerField(),
                "updated": serializers.IntegerField(),
                "errors": serializers.ListField(child=serializers.DictField()),
            },
        ),
    },
    examples=[
        OpenApiExample(
            "Importación exitosa",
            value={"success": True, "data": {"inserted": 40, "updated": 0, "errors": []}, "error": None},
            response_only=True,
        ),
    ],
)
class ImportOtorgamientoView(APIView):
    permission_classes = [IsCareerManager]
    throttle_classes = [BulkOperationRateThrottle]

    def post(self, request):
        if request.user.rol != "jefe_comision" and not request.user.is_superuser:
            return Response({"detail": "Solo el Jefe de Comisión puede importar otorgamientos."}, status=403)
        file = request.FILES.get("file")
        upload_error = validate_excel_upload(file)
        if upload_error:
            return Response({"detail": upload_error}, status=400)
        return _accepted(enqueue_import(request, "otorgamiento", file))


@extend_schema(
    tags=["Importación y exportación"],
    summary="Importar índices de corte por carrera",
    description=(
        "Sube un archivo Excel (.xlsx) con los índices de corte por carrera (Codigo_Carrera, Nombre_Carrera, "
        "Indice_Corte) durante la etapa 6 del proceso de ingreso en curso. Solo puede ejecutarlo el Jefe de "
        "Comisión (o superusuario). La importación es atómica.",
    ),
    request=inline_serializer(
        name="ImportCorteCarreraUploadRequest",
        fields={"file": serializers.FileField(help_text="Archivo .xlsx con los índices de corte.")},
    ),
    responses={
        202: _IMPORT_ACCEPTED,
        201: inline_serializer(
            name="ImportCorteCarreraResponse",
            fields={
                "inserted": serializers.IntegerField(),
                "updated": serializers.IntegerField(),
                "errors": serializers.ListField(child=serializers.DictField(), default=list),
            },
        ),
        400: inline_serializer(
            name="ImportCorteCarreraErrorResponse",
            fields={
                "inserted": serializers.IntegerField(),
                "updated": serializers.IntegerField(),
                "errors": serializers.ListField(child=serializers.DictField()),
            },
        ),
    },
    examples=[
        OpenApiExample(
            "Importación exitosa",
            value={"success": True, "data": {"inserted": 15, "updated": 2, "errors": []}, "error": None},
            response_only=True,
        ),
    ],
)
class ImportCorteCarreraView(APIView):
    permission_classes = [IsCareerManager]
    throttle_classes = [BulkOperationRateThrottle]

    def post(self, request):
        if request.user.rol != "jefe_comision" and not request.user.is_superuser:
            return Response({"detail": "Solo el Jefe de Comisión puede importar índices de corte."}, status=403)
        file = request.FILES.get("file")
        upload_error = validate_excel_upload(file)
        if upload_error:
            return Response({"detail": upload_error}, status=400)
        return _accepted(enqueue_import(request, "corte", file))


@extend_schema(
    tags=["Importación y exportación"],
    summary="Resumen de otorgamientos y cortes",
    description=(
        "Devuelve un resumen numérico de otorgamientos de carreras, estudiantes sin otorgamiento e índices de "
        "corte publicados para un año. Solo accesible al Jefe de Comisión (limitado a su provincia) o superusuario."
    ),
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
    ],
    responses={
        200: inline_serializer(
            name="OtorgamientoSummaryResponse",
            fields={
                "year": serializers.IntegerField(),
                "province": serializers.CharField(),
                "awarded_count": serializers.IntegerField(),
                "without_award_count": serializers.IntegerField(),
                "cuts_count": serializers.IntegerField(),
            },
        ),
    },
    examples=[
        OpenApiExample(
            "Resumen provincial",
            value={
                "success": True,
                "data": {"year": 2025, "province": "La Habana", "awarded_count": 320, "without_award_count": 15, "cuts_count": 48},
                "error": None,
            },
            response_only=True,
        ),
    ],
)
class OtorgamientoSummaryView(APIView):
    permission_classes = [IsCareerManager]

    def get(self, request):
        if request.user.rol != "jefe_comision" and not request.user.is_superuser:
            return Response({"detail": "Solo el Jefe de Comisión puede consultar este resumen."}, status=403)
        try:
            anio = int(request.query_params.get("anio", timezone.now().year))
        except (TypeError, ValueError):
            return Response({"detail": "El año indicado no es válido."}, status=400)
        province_filter = {}
        if request.user.rol == "jefe_comision":
            if not request.user.provincia_id:
                return Response({"detail": "El usuario no tiene una provincia asignada."}, status=400)
            province_filter["estudiante__escuela__municipio__provincia_id"] = request.user.provincia_id
        stage_six_process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[6])
        stage_one_process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[1])
        awards = Otorgamiento.objects.filter(proceso=stage_six_process, **province_filter) if stage_six_process else Otorgamiento.objects.none()
        students = EscalafonItem.objects.filter(
            escalafon__proceso=stage_one_process,
            **{key.replace("estudiante__", "estudiante__"): value for key, value in province_filter.items()},
        ).values("estudiante_id").distinct() if stage_one_process else EscalafonItem.objects.none()
        awarded_student_ids = awards.values("estudiante_id").distinct()
        cuts = CorteCarrera.objects.filter(proceso=stage_six_process) if stage_six_process else CorteCarrera.objects.none()
        if request.user.rol == "jefe_comision":
            provincial_plan_careers = PlanPlaza.objects.filter(
                proceso__anio__year=anio,
                provincia_id=request.user.provincia_id,
            ).values("carrera_id")
            cuts = cuts.filter(carrera_id__in=provincial_plan_careers)
        return Response({
            "year": anio,
            "province": request.user.provincia.nombre if request.user.rol == "jefe_comision" else "Todas",
            "awarded_count": awards.count(),
            "without_award_count": students.exclude(estudiante_id__in=awarded_student_ids).count(),
            "cuts_count": cuts.count(),
        })


@extend_schema(
    tags=["Importación y exportación"],
    summary="Exportar otorgamientos o índices de corte (etapa 6)",
    description=(
        "Exporta a Excel (.xlsx) los otorgamientos de carreras o los índices de corte de la etapa 6 para un año. "
        "El resultado se acota automáticamente al alcance del usuario: Jefe de Comisión a su provincia, "
        "Secretario/Director de escuela a su escuela, superadmin sin restricción."
    ),
    parameters=[
        OpenApiParameter("kind", OpenApiTypes.STR, OpenApiParameter.PATH, enum=["otorgamientos", "cortes"], description="Tipo de exportación."),
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
    ],
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class ExportStageSixView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, kind):
        if request.user.rol not in {"jefe_comision", "secretario_escuela", "director_escuela", "superadmin"} and not request.user.is_superuser:
            return Response({"detail": "No tienes permiso para exportar estos datos."}, status=403)
        try:
            anio = int(request.query_params.get("anio", timezone.now().year))
        except (TypeError, ValueError):
            return Response({"detail": "El año indicado no es válido."}, status=400)
        if kind not in {"otorgamientos", "cortes"}:
            return Response({"detail": "Tipo de exportación no válido."}, status=404)
        process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[6])
        workbook = Workbook()
        sheet = workbook.active
        if kind == "otorgamientos":
            sheet.title = "Otorgamientos"
            sheet.append(["CI", "Código_Carrera", "Nombre_Carrera", "Índice_Otorgamiento"])
            queryset = Otorgamiento.objects.filter(proceso=process) if process else Otorgamiento.objects.none()
            if request.user.rol == "jefe_comision":
                queryset = queryset.filter(estudiante__escuela__municipio__provincia_id=request.user.provincia_id)
            elif request.user.rol in {"secretario_escuela", "director_escuela"}:
                if not request.user.escuela_id:
                    return Response({"detail": "El usuario no tiene una escuela asignada."}, status=400)
                queryset = queryset.filter(estudiante__escuela_id=request.user.escuela_id)
            for item in queryset.select_related("estudiante", "carrera").order_by("estudiante__apellidos", "estudiante__nombre"):
                sheet.append([item.estudiante.ci, item.carrera.codigo, item.carrera.nombre, item.indice_otorgamiento])
        else:
            sheet.title = "Índices de Corte"
            sheet.append(["Código_Carrera", "Nombre_Carrera", "Índice_Corte"])
            queryset = CorteCarrera.objects.filter(proceso=process) if process else CorteCarrera.objects.none()
            if request.user.rol == "jefe_comision":
                career_ids = PlanPlaza.objects.filter(
                    proceso__anio__year=anio,
                    provincia_id=request.user.provincia_id,
                ).values("carrera_id")
                queryset = queryset.filter(carrera_id__in=career_ids)
            elif request.user.rol in {"secretario_escuela", "director_escuela"}:
                if not request.user.escuela_id:
                    return Response({"detail": "El usuario no tiene una escuela asignada."}, status=400)
                career_ids = PlanPlaza.objects.filter(
                    proceso__anio__year=anio,
                    provincia_id=request.user.escuela.municipio.provincia_id,
                ).values("carrera_id")
                queryset = queryset.filter(carrera_id__in=career_ids)
            for item in queryset.select_related("carrera").order_by("carrera__codigo"):
                sheet.append([item.carrera.codigo, item.carrera.nombre, item.indice_corte])
        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{kind}_{anio}.xlsx"'
        return response


@extend_schema(
    tags=["Importación y exportación"],
    summary="Listar otorgamientos de la escuela",
    description=(
        "Lista los otorgamientos de carrera de los estudiantes de la escuela del Secretario o Director de escuela "
        "autenticado, para un año dado, incluyendo el índice general de escalafón de cada estudiante."
    ),
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
    ],
    responses={
        200: inline_serializer(
            name="SchoolOtorgamientoListResponse",
            fields={
                "year": serializers.IntegerField(),
                "items": serializers.ListField(child=inline_serializer(
                    name="SchoolOtorgamientoItem",
                    fields={
                        "id": serializers.IntegerField(),
                        "student": serializers.CharField(),
                        "ci": serializers.CharField(),
                        "general_index": serializers.FloatField(allow_null=True),
                        "award_index": serializers.FloatField(),
                        "career": serializers.CharField(),
                        "ces": serializers.CharField(),
                    },
                )),
            },
        ),
    },
    examples=[
        OpenApiExample(
            "Otorgamientos de la escuela",
            value={
                "success": True,
                "data": {
                    "year": 2025,
                    "items": [{
                        "id": 1, "student": "Ana Pérez", "ci": "01020304050", "general_index": 92.5,
                        "award_index": 90.0, "career": "Medicina", "ces": "Universidad de La Habana",
                    }],
                },
                "error": None,
            },
            response_only=True,
        ),
    ],
)
class SchoolOtorgamientoListView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if request.user.rol not in {"secretario_escuela", "director_escuela"}:
            return Response({"detail": "Solo el Secretario o Director puede consultar estos otorgamientos."}, status=403)
        if not request.user.escuela_id:
            return Response({"detail": "El usuario no tiene una escuela asignada."}, status=400)
        try:
            anio = int(request.query_params.get("anio", timezone.now().year))
        except (TypeError, ValueError):
            return Response({"detail": "El año indicado no es válido."}, status=400)
        process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[6])
        awards = Otorgamiento.objects.filter(
            proceso=process,
            estudiante__escuela_id=request.user.escuela_id,
        ).select_related("estudiante", "carrera__ces") if process else Otorgamiento.objects.none()
        stage_one_process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[1])
        indices = {
            entry.estudiante_id: entry.indice_general
            for entry in EscalafonItem.objects.filter(
                escalafon__proceso=stage_one_process,
                escalafon__escuela_id=request.user.escuela_id,
            )
        } if stage_one_process else {}
        return Response({
            "year": anio,
            "items": [{
                "id": item.id,
                "student": f"{item.estudiante.nombre} {item.estudiante.apellidos}",
                "ci": item.estudiante.ci,
                "general_index": indices.get(item.estudiante_id),
                "award_index": item.indice_otorgamiento,
                "career": item.carrera.nombre,
                "ces": item.carrera.ces.nombre,
            } for item in awards.order_by("estudiante__apellidos", "estudiante__nombre")],
        })


@extend_schema(
    tags=["Importación y exportación"],
    summary="Importar carreras desde Excel",
    description="Sube un archivo Excel (.xlsx) con carreras y las inserta o actualiza en el catálogo. Requiere permiso de gestor de carreras.",
    request=inline_serializer(
        name="ImportCarrerasUploadRequest",
        fields={"file": serializers.FileField(help_text="Archivo .xlsx con las carreras.")},
    ),
    responses={
        202: _IMPORT_ACCEPTED,
        200: inline_serializer(
            name="ImportCarrerasResponse",
            fields={"inserted": serializers.IntegerField(), "errors": serializers.ListField(child=serializers.DictField(), default=list)},
        ),
        400: inline_serializer(
            name="ImportCarrerasErrorResponse",
            fields={"inserted": serializers.IntegerField(), "errors": serializers.ListField(child=serializers.DictField())},
        ),
    },
    examples=[
        OpenApiExample(
            "Importación exitosa",
            value={"success": True, "data": {"inserted": 8, "errors": []}, "error": None},
            response_only=True,
        ),
    ],
)
class ImportCarrerasView(APIView):
    permission_classes = [IsCareerManager]
    throttle_classes = [BulkOperationRateThrottle]

    def post(self, request):
        file = request.FILES.get("file")
        upload_error = validate_excel_upload(file)
        if upload_error:
            return Response({"detail": upload_error}, status=status.HTTP_400_BAD_REQUEST)
        return _accepted(enqueue_import(request, "carreras", file))


@extend_schema(
    tags=["Importación y exportación"],
    summary="Exportar carreras a Excel",
    description="Exporta a Excel (.xlsx) el catálogo completo de carreras. Requiere permiso de gestor de carreras.",
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class ExportCarrerasView(APIView):
    permission_classes = [IsCareerManager]

    def get(self, request):
        data = CareerExcelService().export_file(Carrera.objects.all().order_by("nombre"))
        response = HttpResponse(
            data,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="carreras.xlsx"'
        return response


@extend_schema(
    tags=["Importación y exportación"],
    summary="Exportar catálogo genérico",
    description=(
        "Exporta a Excel (.xlsx) un catálogo del sistema (Provincia, Municipio, Escuela, Ces, Carrera, Asignatura "
        "o TipoOtorgamiento, según la subclase de vista concreta registrada en la URL). Incluye id, nombre, "
        "descripción, estado y fecha de última modificación."
    ),
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class ExportCatalogView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    resource_model = None
    filename = "catalogo.xlsx"

    def get(self, request, *args, **kwargs):
        if not self.resource_model:
            return Response({"detail": "Debe indicar un modelo de catálogo."}, status=400)

        model_map = {
            "Provincia": Provincia,
            "Municipio": Municipio,
            "Escuela": Escuela,
            "Ces": Ces,
            "Carrera": Carrera,
            "Asignatura": Asignatura,
            "TipoOtorgamiento": TipoOtorgamiento,
        }
        model = model_map.get(self.resource_model)
        if model is None:
            return Response({"detail": "Modelo de catálogo no soportado."}, status=400)

        queryset = model.objects.all().order_by("nombre")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.resource_model
        headers = ["id", "nombre", "descripcion", "estado", "fecha_ultima_modificacion"]
        sheet.append(headers)
        for item in queryset:
            estado = getattr(item, "activa", None)
            if estado is None:
                estado = getattr(item, "activo", None)
            last_modified = getattr(item, "fecha_ultima_modificacion", None)
            sheet.append([
                item.pk,
                getattr(item, "nombre", ""),
                getattr(item, "descripcion", "") or "",
                "Activo" if estado else "Inactivo",
                timezone.localtime(last_modified).strftime("%Y-%m-%d %H:%M:%S") if last_modified else "",
            ])
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{self.filename}"'
        return response


@extend_schema(
    tags=["Escalafón"],
    summary="Importar escalafón de una escuela",
    description=(
        "Sube un archivo Excel (.xlsx) con el escalafón (índices de 10mo, 11no, 12mo y general) de los estudiantes "
        "de una escuela para un año/proceso dado. Solo pueden importarlo el Secretario de la escuela (limitado a "
        "su propia escuela), el Jefe de Comisión (limitado a escuelas de su provincia) o el superadmin. Falla si "
        "ya existe un escalafón para esa escuela y año. Al finalizar sin errores, notifica a los estudiantes de la "
        "escuela.",
    ),
    request=inline_serializer(
        name="ImportEscalafonUploadRequest",
        fields={
            "file": serializers.FileField(help_text="Archivo .xlsx con el escalafón."),
            "escuela": serializers.CharField(required=False, help_text="ID de la escuela (requerido salvo para Secretario, que usa la suya)."),
            "anio": serializers.IntegerField(required=False, help_text="Año del proceso de escalafón (también acepta 'año'). Por defecto el año actual."),
        },
    ),
    responses={
        202: _IMPORT_ACCEPTED,
        201: inline_serializer(
            name="ImportEscalafonResponse",
            fields={"inserted": serializers.IntegerField(), "errors": serializers.ListField(child=serializers.DictField(), default=list)},
        ),
        400: inline_serializer(
            name="ImportEscalafonErrorResponse",
            fields={"inserted": serializers.IntegerField(), "errors": serializers.ListField(child=serializers.DictField())},
        ),
    },
    examples=[
        OpenApiExample(
            "Importación exitosa",
            value={"success": True, "data": {"inserted": 30, "errors": []}, "error": None},
            response_only=True,
        ),
    ],
)
class ImportEscalafonView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    throttle_classes = [BulkOperationRateThrottle]

    def post(self, request):
        if request.user.rol not in {"secretario_escuela", "jefe_comision", "superadmin"}:
            return Response({"detail": "Solo el Secretario puede importar el escalafón."}, status=status.HTTP_403_FORBIDDEN)
        if request.user.rol == "secretario_escuela" and request.data.get("escuela") not in {None, "", str(request.user.escuela_id), request.user.escuela_id}:
            return Response({"detail": "Solo puedes importar el escalafón de tu escuela."}, status=status.HTTP_403_FORBIDDEN)
        file = request.FILES.get("file")
        upload_error = validate_excel_upload(file)
        if upload_error:
            return Response({"detail": upload_error}, status=status.HTTP_400_BAD_REQUEST)
        try:
            escuela = resolve_school(request.data.get("escuela") or request.user.escuela_id)
            anio = int(request.data.get("anio", request.data.get("año", timezone.now().year)))
            proceso = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[1])
            if not proceso:
                raise ValueError("No existe un proceso para la etapa de escalafón en el año indicado.")
        except (TypeError, ValueError, Escuela.DoesNotExist):
            return Response({"detail": "Debe indicar una escuela válida y el año del escalafón."}, status=status.HTTP_400_BAD_REQUEST)
        if (
            request.user.rol == "jefe_comision"
            and escuela.municipio.provincia_id != request.user.provincia_id
        ):
            return Response({"detail": "Solo puedes importar el escalafón de escuelas de tu provincia."}, status=status.HTTP_403_FORBIDDEN)
        task_id = enqueue_import(request, "escalafon", file, {"escuela_id": escuela.id, "proceso_id": proceso.id, "anio": anio})
        return _accepted(task_id)


@extend_schema(
    tags=["Resultados"],
    summary="Importar resultados de exámenes",
    description=(
        "Sube un archivo Excel (.xlsx) con las notas de una asignatura para el proceso de resultados (etapa 5) de "
        "un año dado. Solo el Jefe de Comisión (limitado a su provincia) o el superadmin pueden importar, y solo "
        "mientras la etapa 5 esté 'en_curso'. Al finalizar sin errores, notifica a los estudiantes afectados y "
        "registra auditoría.",
    ),
    request=inline_serializer(
        name="ImportResultadosUploadRequest",
        fields={
            "file": serializers.FileField(help_text="Archivo .xlsx con las notas."),
            "anio": serializers.IntegerField(required=False, help_text="Año del proceso de resultados. Por defecto el año actual."),
            "asignatura": serializers.CharField(help_text="Nombre de la asignatura importada (Matemática, Español o Historia)."),
            "fecha_limite_reclamo": serializers.DateField(help_text="Fecha límite (ISO 8601, YYYY-MM-DD) para presentar reclamaciones sobre estas notas."),
        },
    ),
    responses={
        202: _IMPORT_ACCEPTED,
        201: inline_serializer(
            name="ImportResultadosResponse",
            fields={
                "inserted": serializers.IntegerField(),
                "updated": serializers.IntegerField(),
                "errors": serializers.ListField(child=serializers.DictField(), default=list),
            },
        ),
        400: inline_serializer(
            name="ImportResultadosErrorResponse",
            fields={
                "inserted": serializers.IntegerField(),
                "updated": serializers.IntegerField(),
                "errors": serializers.ListField(child=serializers.DictField()),
            },
        ),
    },
    examples=[
        OpenApiExample(
            "Importación exitosa",
            value={"success": True, "data": {"inserted": 100, "updated": 0, "errors": []}, "error": None},
            response_only=True,
        ),
    ],
)
class ImportResultadosView(APIView):
    permission_classes = [permissions.IsAuthenticated]
    throttle_classes = [BulkOperationRateThrottle]

    def post(self, request):
        if request.user.rol not in {"jefe_comision", "superadmin"}:
            return Response({"detail": "Solo el Jefe de Comisión puede importar resultados."}, status=403)
        if not request.user.provincia_id and request.user.rol != "superadmin":
            return Response({"detail": "El usuario no tiene una provincia asignada."}, status=400)
        file = request.FILES.get("file")
        upload_error = validate_excel_upload(file)
        if upload_error:
            return Response({"detail": upload_error}, status=400)
        try:
            anio = int(request.data.get("anio", timezone.now().year))
            proceso = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[5])
        except (TypeError, ValueError):
            proceso = None
        if not proceso:
            return Response({"detail": "No existe un proceso de resultados para el año indicado."}, status=400)
        if proceso.etapa.estado != "en_curso":
            return Response({"detail": "La importación de resultados solo está disponible durante la etapa 5."}, status=403)
        selected_subject = request.data.get("asignatura")
        try:
            deadline = date.fromisoformat(str(request.data.get("fecha_limite_reclamo", "")))
        except ValueError:
            return Response({"detail": "Debe indicar una fecha límite de reclamaciones válida."}, status=400)
        task_id = enqueue_import(request, "resultados", file, {
            "proceso_id": proceso.id, "asignatura": selected_subject, "deadline": deadline.isoformat(),
        })
        return _accepted(task_id)


def normalize_result_subject(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    if "matematic" in text:
        return "matematica"
    if "espanol" in text:
        return "espanol"
    if "historia" in text:
        return "historia"
    return None


def _school_result_row(entry, grades, anio):
    return {
        "id": entry.id,
        "ci": entry.estudiante.ci,
        "student": f"{entry.estudiante.nombre} {entry.estudiante.apellidos}",
        "school": entry.escalafon.escuela.nombre,
        "matematica": grades.get("matematica"),
        "espanol": grades.get("espanol"),
        "historia": grades.get("historia"),
        "escalafon_index": float(entry.indice_general),
        "process_year": anio,
    }


@extend_schema(
    tags=["Resultados"],
    summary="Consultar resultados de exámenes",
    description=(
        "Consulta los resultados de exámenes de ingreso para un año. La forma de la respuesta depende del rol del "
        "usuario autenticado: Secretario/Director de escuela recibe la lista de estudiantes de su escuela con sus "
        "notas; un estudiante recibe el estado de la etapa y sus propias notas/reclamaciones por asignatura; "
        "Jefe de Comisión (acotado a su provincia) o superadmin reciben la lista completa de resultados."
    ),
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
    ],
    responses={200: OpenApiTypes.OBJECT},
    examples=[
        OpenApiExample(
            "Vista de Secretario/Director de escuela",
            value={"success": True, "data": [{
                "id": 1, "ci": "01020304050", "student": "Ana Pérez", "school": "IPU José Martí",
                "matematica": 85, "espanol": 90, "historia": 88, "escalafon_index": 92.5, "process_year": 2025,
            }], "error": None},
            response_only=True,
        ),
        OpenApiExample(
            "Vista de estudiante",
            value={"success": True, "data": {
                "stage": {"active": True, "completed": False, "current_number": 5, "fecha_fin": "2025-07-15"},
                "results": [{
                    "id": 10, "subject": "Matemática", "grade": 85, "fecha_limite_reclamo": "2025-07-20",
                    "claim": None,
                }],
            }, "error": None},
            response_only=True,
        ),
        OpenApiExample(
            "Vista de Jefe de Comisión/superadmin",
            value={"success": True, "data": [{
                "id": 1, "ci": "01020304050", "student": "Ana Pérez", "subject": "Matemática", "grade": 85,
                "school": "IPU José Martí", "fecha_limite_reclamo": "2025-07-20", "process_year": 2025,
            }], "error": None},
            response_only=True,
        ),
    ],
)
class ResultadosListView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        anio = int(request.query_params.get("anio", timezone.now().year))
        if request.user.rol in {"secretario_escuela", "director_escuela"}:
            if not request.user.escuela_id:
                return Response({"detail": "El usuario no tiene una escuela asignada."}, status=400)
            escalafon_process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[1])
            if not escalafon_process:
                return Response([])
            entries = EscalafonItem.objects.filter(
                escalafon__proceso=escalafon_process,
                escalafon__escuela_id=request.user.escuela_id,
            ).select_related("estudiante", "escalafon__escuela")
            result_process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[5])
            results = ResultadoExamen.objects.filter(
                proceso=result_process,
                estudiante__escuela_id=request.user.escuela_id,
            ).select_related("asignatura") if result_process else ResultadoExamen.objects.none()
            results_by_student = {}
            for result in results:
                subject_key = normalize_result_subject(result.asignatura.nombre)
                if subject_key:
                    results_by_student.setdefault(result.estudiante_id, {})[subject_key] = result.nota
            return Response([_school_result_row(entry, results_by_student.get(entry.estudiante_id, {}), anio) for entry in entries.order_by("-indice_general", "estudiante__apellidos")])

        if request.user.rol == "estudiante":
            student = getattr(request.user, "estudiante", None)
            if not student:
                return Response({"detail": "El usuario no tiene un perfil de estudiante."}, status=400)
            result_process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[5])
            results = ResultadoExamen.objects.filter(
                estudiante=student,
                proceso=result_process,
            ).select_related("asignatura") if result_process else ResultadoExamen.objects.none()
            claims = {
                claim.resultado_id: claim
                for claim in Reclamacion.objects.filter(
                    estudiante=student,
                    resultado__proceso=result_process,
                )
            } if result_process else {}
            current_stage = Etapa.objects.filter(estado="en_curso").order_by("id").first()
            current_stage_number = next(
                (number for number, name in ETAPAS_NOMBRES.items() if current_stage and current_stage.nombre == name),
                None,
            )
            results_by_subject = {
                normalize_result_subject(result.asignatura.nombre): result
                for result in results
                if normalize_result_subject(result.asignatura.nombre)
            }
            subjects = [("matematica", "Matemática"), ("espanol", "Español"), ("historia", "Historia")]
            return Response({
                "stage": {
                    "active": current_stage_number == 5,
                    "completed": current_stage_number is not None and current_stage_number > 5,
                    "current_number": current_stage_number,
                    "fecha_fin": current_stage.fecha_fin if current_stage_number == 5 else None,
                },
                "results": [{
                    "id": result.id if result else subject_key,
                    "subject": subject_name,
                    "grade": result.nota if result else None,
                    "fecha_limite_reclamo": result.fecha_limite_reclamo if result else None,
                    "claim": (
                        {
                            "id": claims[result.id].id,
                            "status": claims[result.id].estado,
                            "description": claims[result.id].descripcion,
                            "fecha_presentacion": claims[result.id].fecha_presentacion,
                            "lugar_presentacion": claims[result.id].lugar_presentacion,
                        }
                        if result and result.id in claims else None
                    ),
                } for subject_key, subject_name in subjects for result in [results_by_subject.get(subject_key)]],
            })
        elif request.user.rol in {"jefe_comision", "superadmin"}:
            queryset = ResultadoExamen.objects.filter(proceso__etapa__nombre=ETAPAS_NOMBRES[5])
            if request.user.rol == "jefe_comision":
                queryset = queryset.filter(estudiante__escuela__municipio__provincia_id=request.user.provincia_id)
        else:
            return Response({"detail": "No tienes permiso para consultar resultados."}, status=403)
        return Response([{
            "id": item.id,
            "ci": item.estudiante.ci,
            "student": f"{item.estudiante.nombre} {item.estudiante.apellidos}",
            "subject": item.asignatura.nombre,
            "grade": item.nota,
            "school": item.estudiante.escuela.nombre,
            "fecha_limite_reclamo": item.fecha_limite_reclamo,
            "puede_reclamar": not (item.fecha_limite_reclamo and item.fecha_limite_reclamo < timezone.localdate()),
            "process_year": item.proceso.anio.year,
        } for item in queryset.select_related("estudiante__escuela", "asignatura", "proceso").order_by("estudiante__apellidos", "asignatura__nombre")])


@extend_schema(
    tags=["Resultados"],
    summary="Resultados públicos",
    description=(
        "Consulta pública (sin autenticación) de resultados de exámenes de ingreso para la interoperabilidad "
        "institucional, con filtros por año, provincia, municipio, escuela y asignatura. También devuelve los "
        "valores disponibles para cada filtro."
    ),
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
        OpenApiParameter("provincia", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre exacto de la provincia."),
        OpenApiParameter("municipio", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre exacto del municipio."),
        OpenApiParameter("escuela", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre exacto de la escuela."),
        OpenApiParameter("asignatura", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre exacto de la asignatura."),
    ],
    responses={
        200: inline_serializer(
            name="LandingResultadosResponse",
            fields={
                "year": serializers.IntegerField(),
                "results": serializers.ListField(child=inline_serializer(
                    name="LandingResultadoItem",
                    fields={
                        "id": serializers.IntegerField(),
                        "student": serializers.CharField(),
                        "school": serializers.CharField(),
                        "municipality": serializers.CharField(),
                        "province": serializers.CharField(),
                        "subject": serializers.CharField(),
                        "grade": serializers.FloatField(),
                        "year": serializers.IntegerField(),
                    },
                )),
                "years": serializers.ListField(child=serializers.IntegerField()),
                "provinces": serializers.ListField(child=serializers.CharField()),
                "municipalities": serializers.ListField(child=serializers.CharField()),
                "schools": serializers.ListField(child=serializers.CharField()),
                "subjects": serializers.ListField(child=serializers.CharField()),
            },
        ),
        400: inline_serializer(name="LandingResultadosErrorResponse", fields={"detail": serializers.CharField()}),
    },
    examples=[
        OpenApiExample(
            "Resultados filtrados",
            value={
                "success": True,
                "data": {
                    "year": 2025,
                    "results": [{
                        "id": 1, "student": "Ana Pérez", "school": "IPU José Martí", "municipality": "Plaza",
                        "province": "La Habana", "subject": "Matemática", "grade": 85, "year": 2025,
                    }],
                    "years": [2025, 2024],
                    "provinces": ["La Habana"],
                    "municipalities": ["Plaza"],
                    "schools": ["IPU José Martí"],
                    "subjects": ["Matemática", "Español", "Historia"],
                },
                "error": None,
            },
            response_only=True,
        ),
    ],
)
class LandingResultadosView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        try:
            anio = int(request.query_params.get("anio", timezone.now().year))
        except (TypeError, ValueError):
            return Response({"detail": "El año indicado no es válido."}, status=400)
        queryset = ResultadoExamen.objects.filter(
            proceso__anio__year=anio,
            proceso__etapa__nombre=ETAPAS_NOMBRES[5],
        ).select_related("estudiante__escuela__municipio__provincia", "asignatura", "proceso")
        provincia = request.query_params.get("provincia", "").strip()
        municipio = request.query_params.get("municipio", "").strip()
        escuela = request.query_params.get("escuela", "").strip()
        asignatura = request.query_params.get("asignatura", "").strip()
        if provincia:
            queryset = queryset.filter(estudiante__escuela__municipio__provincia__nombre__iexact=provincia)
        if municipio:
            queryset = queryset.filter(estudiante__escuela__municipio__nombre__iexact=municipio)
        if escuela:
            queryset = queryset.filter(estudiante__escuela__nombre__iexact=escuela)
        if asignatura:
            queryset = queryset.filter(asignatura__nombre__iexact=asignatura)
        results = [{
            "id": item.id,
            "student": f"{item.estudiante.nombre} {item.estudiante.apellidos}",
            "school": item.estudiante.escuela.nombre,
            "municipality": item.estudiante.escuela.municipio.nombre,
            "province": item.estudiante.escuela.municipio.provincia.nombre,
            "subject": item.asignatura.nombre,
            "grade": item.nota,
            "year": item.proceso.anio.year,
        } for item in queryset.order_by("estudiante__apellidos", "estudiante__nombre", "asignatura__nombre")]
        all_results = ResultadoExamen.objects.filter(
            proceso__anio__year=anio,
            proceso__etapa__nombre=ETAPAS_NOMBRES[5],
        ).select_related("estudiante__escuela__municipio__provincia", "asignatura")
        return Response({
            "year": anio,
            "results": results,
            "years": list(Proceso.objects.filter(etapa__nombre=ETAPAS_NOMBRES[5]).values_list("anio__year", flat=True).distinct().order_by("-anio__year")),
            "provinces": list(all_results.values_list("estudiante__escuela__municipio__provincia__nombre", flat=True).distinct().order_by("estudiante__escuela__municipio__provincia__nombre")),
            "municipalities": list(all_results.values_list("estudiante__escuela__municipio__nombre", flat=True).distinct().order_by("estudiante__escuela__municipio__nombre")),
            "schools": list(all_results.values_list("estudiante__escuela__nombre", flat=True).distinct().order_by("estudiante__escuela__nombre")),
            "subjects": list(all_results.values_list("asignatura__nombre", flat=True).distinct().order_by("asignatura__nombre")),
        })


@extend_schema(
    tags=["Otorgamiento"],
    summary="Otorgamientos públicos",
    description=(
        "Consulta pública (sin autenticación) del resultado de otorgamiento de carreras, con filtros por año, "
        "provincia y CI, para la interoperabilidad institucional. También devuelve los años y provincias "
        "disponibles."
    ),
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
        OpenApiParameter("provincia", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre exacto de la provincia."),
        OpenApiParameter("ci", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Carné de identidad (búsqueda parcial) del estudiante."),
    ],
    responses={
        200: inline_serializer(
            name="LandingOtorgamientosResponse",
            fields={
                "year": serializers.IntegerField(),
                "results": serializers.ListField(child=inline_serializer(
                    name="LandingOtorgamientoItem",
                    fields={
                        "id": serializers.IntegerField(),
                        "student": serializers.CharField(),
                        "ci": serializers.CharField(),
                        "career": serializers.CharField(),
                        "ces": serializers.CharField(),
                        "award_index": serializers.FloatField(),
                        "school": serializers.CharField(),
                        "province": serializers.CharField(),
                        "year": serializers.IntegerField(),
                    },
                )),
                "years": serializers.ListField(child=serializers.IntegerField()),
                "provinces": serializers.ListField(child=serializers.CharField()),
            },
        ),
        400: inline_serializer(name="LandingOtorgamientosErrorResponse", fields={"detail": serializers.CharField()}),
    },
    examples=[
        OpenApiExample(
            "Otorgamientos filtrados",
            value={
                "success": True,
                "data": {
                    "year": 2025,
                    "results": [{
                        "id": 1, "student": "Ana Pérez", "ci": "01020304050", "career": "Medicina",
                        "ces": "Universidad de La Habana", "award_index": 90.0, "school": "IPU José Martí",
                        "province": "La Habana", "year": 2025,
                    }],
                    "years": [2025, 2024],
                    "provinces": ["La Habana"],
                },
                "error": None,
            },
            response_only=True,
        ),
    ],
)
class LandingOtorgamientosView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        try:
            anio = int(request.query_params.get("anio", timezone.now().year))
        except (TypeError, ValueError):
            return Response({"detail": "El año indicado no es válido."}, status=400)
        provincia = request.query_params.get("provincia", "").strip()
        ci = request.query_params.get("ci", "").strip()
        queryset = Otorgamiento.objects.filter(
            proceso__anio__year=anio,
            proceso__etapa__nombre=ETAPAS_NOMBRES[6],
        ).select_related("estudiante__escuela__municipio__provincia", "carrera__ces", "proceso")
        if provincia:
            queryset = queryset.filter(estudiante__escuela__municipio__provincia__nombre__iexact=provincia)
        if ci:
            queryset = queryset.filter(estudiante__ci__icontains=ci)
        all_awards = Otorgamiento.objects.filter(
            proceso__anio__year=anio,
            proceso__etapa__nombre=ETAPAS_NOMBRES[6],
        ).select_related("estudiante__escuela__municipio__provincia")
        return Response({
            "year": anio,
            "results": [{
                "id": item.id,
                "student": f"{item.estudiante.nombre} {item.estudiante.apellidos}",
                "ci": item.estudiante.ci,
                "career": item.carrera.nombre,
                "ces": item.carrera.ces.nombre,
                "award_index": item.indice_otorgamiento,
                "school": item.estudiante.escuela.nombre,
                "province": item.estudiante.escuela.municipio.provincia.nombre,
                "year": item.proceso.anio.year,
            } for item in queryset.order_by("estudiante__apellidos", "estudiante__nombre")],
            "years": list(Proceso.objects.filter(etapa__nombre=ETAPAS_NOMBRES[6]).values_list("anio__year", flat=True).distinct().order_by("-anio__year")),
            "provinces": list(all_awards.values_list("estudiante__escuela__municipio__provincia__nombre", flat=True).distinct().order_by("estudiante__escuela__municipio__provincia__nombre")),
        })


@extend_schema(
    tags=["Resultados"],
    summary="Índices de corte públicos",
    description=(
        "Consulta pública (sin autenticación) de los índices de corte por carrera del año más reciente disponible "
        "(o el indicado), con filtros por provincia y nombre de carrera, ordenados por cantidad de solicitudes "
        "recibidas y nombre de carrera."
    ),
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año más reciente con índices de corte publicados."),
        OpenApiParameter("provincia", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre exacto de la provincia de la universidad."),
        OpenApiParameter("carrera", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Búsqueda parcial por nombre de carrera."),
    ],
    responses={
        200: inline_serializer(
            name="LandingCortesResponse",
            fields={
                "year": serializers.IntegerField(allow_null=True),
                "years": serializers.ListField(child=serializers.IntegerField()),
                "provinces": serializers.ListField(child=serializers.CharField()),
                "items": serializers.ListField(child=inline_serializer(
                    name="LandingCorteItem",
                    fields={
                        "id": serializers.IntegerField(),
                        "career": serializers.CharField(),
                        "career_code": serializers.CharField(),
                        "index": serializers.FloatField(),
                        "requests_count": serializers.IntegerField(),
                        "year": serializers.IntegerField(),
                    },
                )),
            },
        ),
        400: inline_serializer(name="LandingCortesErrorResponse", fields={"detail": serializers.CharField()}),
    },
    examples=[
        OpenApiExample(
            "Índices de corte",
            value={
                "success": True,
                "data": {
                    "year": 2025,
                    "years": [2025, 2024],
                    "provinces": ["La Habana"],
                    "items": [{"id": 1, "career": "Medicina", "career_code": "MED01", "index": 88.5, "requests_count": 120, "year": 2025}],
                },
                "error": None,
            },
            response_only=True,
        ),
    ],
)
class LandingCortesView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request):
        available_years = list(CorteCarrera.objects.filter(
            proceso__etapa__nombre=ETAPAS_NOMBRES[6],
        ).values_list("proceso__anio__year", flat=True).distinct().order_by("-proceso__anio__year"))
        if not available_years:
            return Response({"year": None, "items": [], "years": []})
        try:
            selected_year = int(request.query_params.get("anio", available_years[0]))
        except (TypeError, ValueError):
            return Response({"detail": "El año indicado no es válido."}, status=400)
        if selected_year not in available_years:
            selected_year = available_years[0]
        cuts = CorteCarrera.objects.filter(
            proceso__anio__year=selected_year,
            proceso__etapa__nombre=ETAPAS_NOMBRES[6],
        ).select_related("carrera", "proceso")
        province = request.query_params.get("provincia", "").strip()
        career_search = request.query_params.get("carrera", "").strip()
        if province:
            province_careers = PlanPlaza.objects.filter(
                proceso__anio__year=selected_year,
                provincia__nombre__iexact=province,
            ).values("carrera_id")
            cuts = cuts.filter(carrera_id__in=province_careers)
        if career_search:
            cuts = cuts.filter(carrera__nombre__icontains=career_search)
        requests = BoletaSolicitudItem.objects.filter(
            boleta_solicitud__proceso__anio__year=selected_year,
            boleta_solicitud__proceso__etapa__nombre=ETAPAS_NOMBRES[3],
        ).values("plan_plaza__carrera_id").annotate(total=Count("id"))
        requested_counts = {row["plan_plaza__carrera_id"]: row["total"] for row in requests}
        cuts = sorted(cuts, key=lambda cut: (-requested_counts.get(cut.carrera_id, 0), cut.carrera.nombre))
        return Response({
            "year": selected_year,
            "years": available_years,
            "provinces": list(PlanPlaza.objects.filter(proceso__anio__year=selected_year).values_list("provincia__nombre", flat=True).distinct().order_by("provincia__nombre")),
            "items": [{
                "id": cut.id,
                "career": cut.carrera.nombre,
                "career_code": cut.carrera.codigo,
                "index": cut.indice_corte,
                "requests_count": requested_counts.get(cut.carrera_id, 0),
                "year": selected_year,
            } for cut in cuts],
        })


@extend_schema(
    tags=["Importación y exportación"],
    summary="Exportar datos públicos a Excel",
    description=(
        "Exporta a Excel (.xlsx), sin autenticación, resultados de exámenes, otorgamientos de carreras o índices "
        "de corte, filtrables por año y provincia, para la interoperabilidad institucional."
    ),
    parameters=[
        OpenApiParameter("kind", OpenApiTypes.STR, OpenApiParameter.PATH, enum=["resultados", "otorgamientos", "cortes"], description="Tipo de exportación."),
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
        OpenApiParameter("provincia", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre exacto de la provincia para filtrar."),
    ],
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class LandingExcelExportView(APIView):
    permission_classes = [permissions.AllowAny]

    def get(self, request, kind):
        if kind not in {"resultados", "otorgamientos", "cortes"}:
            return Response({"detail": "Tipo de exportación no válido."}, status=404)
        try:
            anio = int(request.query_params.get("anio", timezone.now().year))
        except (TypeError, ValueError):
            return Response({"detail": "El año indicado no es válido."}, status=400)
        provincia = request.query_params.get("provincia", "").strip()
        workbook = Workbook()
        sheet = workbook.active

        if kind == "resultados":
            sheet.title = "Notas de ingreso"
            sheet.append(["CI", "Estudiante", "Asignatura", "Nota", "Escuela", "Municipio", "Provincia", "Año"])
            queryset = ResultadoExamen.objects.filter(
                proceso__anio__year=anio,
                proceso__etapa__nombre=ETAPAS_NOMBRES[5],
            ).select_related("estudiante__escuela__municipio__provincia", "asignatura")
            if provincia:
                queryset = queryset.filter(estudiante__escuela__municipio__provincia__nombre__iexact=provincia)
            for item in queryset.order_by("estudiante__apellidos", "estudiante__nombre", "asignatura__nombre"):
                sheet.append([
                    item.estudiante.ci,
                    f"{item.estudiante.nombre} {item.estudiante.apellidos}",
                    item.asignatura.nombre,
                    item.nota,
                    item.estudiante.escuela.nombre,
                    item.estudiante.escuela.municipio.nombre,
                    item.estudiante.escuela.municipio.provincia.nombre,
                    anio,
                ])
        elif kind == "otorgamientos":
            sheet.title = "Otorgamientos"
            sheet.append(["CI", "Estudiante", "Carrera", "CES", "Índice", "Escuela", "Provincia", "Año"])
            queryset = Otorgamiento.objects.filter(
                proceso__anio__year=anio,
                proceso__etapa__nombre=ETAPAS_NOMBRES[6],
            ).select_related("estudiante__escuela__municipio__provincia", "carrera__ces")
            if provincia:
                queryset = queryset.filter(estudiante__escuela__municipio__provincia__nombre__iexact=provincia)
            for item in queryset.order_by("estudiante__apellidos", "estudiante__nombre"):
                sheet.append([
                    item.estudiante.ci,
                    f"{item.estudiante.nombre} {item.estudiante.apellidos}",
                    item.carrera.nombre,
                    item.carrera.ces.nombre,
                    item.indice_otorgamiento,
                    item.estudiante.escuela.nombre,
                    item.estudiante.escuela.municipio.provincia.nombre,
                    anio,
                ])
        else:
            sheet.title = "Índices de corte"
            sheet.append(["Código", "Carrera", "Índice de corte", "Provincia", "Año"])
            queryset = CorteCarrera.objects.filter(
                proceso__anio__year=anio,
                proceso__etapa__nombre=ETAPAS_NOMBRES[6],
            ).select_related("carrera", "proceso")
            if provincia:
                province_careers = PlanPlaza.objects.filter(
                    proceso__anio__year=anio,
                    provincia__nombre__iexact=provincia,
                ).values("carrera_id")
                queryset = queryset.filter(carrera_id__in=province_careers)
            province_names = dict(
                PlanPlaza.objects.filter(proceso__anio__year=anio)
                .values_list("carrera_id", "provincia__nombre")
            )
            for item in queryset.order_by("carrera__codigo"):
                sheet.append([item.carrera.codigo, item.carrera.nombre, item.indice_corte, province_names.get(item.carrera_id, ""), anio])

        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{kind}_{provincia or "todas"}_{anio}.xlsx"'
        return response


@extend_schema(
    tags=["Resultados"],
    summary="Presentar reclamación sobre una nota",
    description=(
        "Permite a un estudiante autenticado presentar una reclamación sobre una de sus notas publicadas, "
        "mientras la etapa 5 esté 'en_curso' y no haya vencido la fecha límite de reclamación. Solo se admite una "
        "reclamación por resultado."
    ),
    parameters=[
        OpenApiParameter("result_id", OpenApiTypes.INT, OpenApiParameter.PATH, description="ID del resultado de examen a reclamar."),
    ],
    request=inline_serializer(
        name="StudentResultClaimRequest",
        fields={"descripcion": serializers.CharField(max_length=500, help_text="Motivo de la reclamación (máximo 500 caracteres).")},
    ),
    responses={
        201: inline_serializer(name="StudentResultClaimResponse", fields={"id": serializers.IntegerField(), "status": serializers.CharField()}),
        400: inline_serializer(name="StudentResultClaimErrorResponse", fields={"detail": serializers.CharField()}),
    },
    examples=[
        OpenApiExample("Reclamación creada", value={"success": True, "data": {"id": 5, "status": "pendiente"}, "error": None}, response_only=True),
    ],
)
class StudentResultClaimView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, result_id):
        if request.user.rol != "estudiante":
            return Response({"detail": "Solo un estudiante puede reclamar una nota."}, status=403)
        result = ResultadoExamen.objects.filter(
            pk=result_id,
            estudiante__usuario=request.user,
            proceso__etapa__nombre=ETAPAS_NOMBRES[5],
        ).first()
        if not result:
            return Response({"detail": "No existe ese resultado publicado."}, status=404)
        stage = Etapa.objects.filter(nombre=ETAPAS_NOMBRES[5]).first()
        if not stage or stage.estado != "en_curso":
            return Response({"detail": "Las reclamaciones no están disponibles fuera de la etapa 5."}, status=403)
        if result.fecha_limite_reclamo and result.fecha_limite_reclamo < timezone.localdate():
            return Response({"detail": "El plazo de reclamación ya finalizó."}, status=403)
        if Reclamacion.objects.filter(resultado=result).exists():
            return Response({"detail": "Ya existe una reclamación para esta nota."}, status=400)
        description = str(request.data.get("descripcion", "")).strip()
        if not description:
            return Response({"detail": "Debes indicar el motivo de la reclamación."}, status=400)
        if len(description) > 500:
            return Response({"detail": "El motivo de la reclamación no puede superar los 500 caracteres."}, status=400)
        claim = Reclamacion.objects.create(
            estudiante=request.user.estudiante,
            resultado=result,
            descripcion=description,
        )
        record_audit(
            request.user,
            "Presentación de reclamo de nota",
            request.path,
            request=request,
            new={"reclamacion_id": claim.id, "resultado_id": result.id, "asignatura": result.asignatura.nombre},
        )
        return Response({"id": claim.id, "status": claim.estado}, status=201)


@extend_schema(
    tags=["Resultados"],
    summary="Listar reclamaciones pendientes",
    description=(
        "Lista las reclamaciones de notas pendientes de resolución para un año, acotadas a la provincia del Jefe "
        "de Comisión, o sin restricción para el superadmin."
    ),
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
    ],
    responses={
        200: inline_serializer(
            name="ResultClaimsListResponse",
            fields={
                "id": serializers.IntegerField(),
                "student": serializers.CharField(),
                "ci": serializers.CharField(),
                "school": serializers.CharField(),
                "subject": serializers.CharField(),
                "grade": serializers.FloatField(),
                "description": serializers.CharField(),
                "status": serializers.CharField(),
                "date": serializers.DateField(),
                "deadline": serializers.DateField(allow_null=True),
                "process_year": serializers.IntegerField(),
            },
            many=True,
        ),
    },
    examples=[
        OpenApiExample(
            "Reclamaciones pendientes",
            value={"success": True, "data": [{
                "id": 5, "student": "Ana Pérez", "ci": "01020304050", "school": "IPU José Martí",
                "subject": "Matemática", "grade": 60, "description": "Solicito revisión de la nota.",
                "status": "pendiente", "date": "2025-07-01", "deadline": "2025-07-20", "process_year": 2025,
            }], "error": None},
            response_only=True,
        ),
    ],
)
class ResultClaimsListView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if request.user.rol not in {"jefe_comision", "superadmin"}:
            return Response({"detail": "No tienes permiso para consultar reclamaciones."}, status=403)
        try:
            anio = int(request.query_params.get("anio", timezone.now().year))
        except (TypeError, ValueError):
            return Response({"detail": "El año indicado no es válido."}, status=400)
        claims = Reclamacion.objects.filter(
            estado="pendiente",
            resultado__proceso__anio__year=anio,
            resultado__proceso__etapa__nombre=ETAPAS_NOMBRES[5],
        ).select_related("estudiante__escuela", "resultado__asignatura", "resultado__proceso")
        if request.user.rol == "jefe_comision":
            claims = claims.filter(estudiante__escuela__municipio__provincia_id=request.user.provincia_id)
        return Response([{
            "id": claim.id,
            "student": f"{claim.estudiante.nombre} {claim.estudiante.apellidos}",
            "ci": claim.estudiante.ci,
            "school": claim.estudiante.escuela.nombre,
            "subject": claim.resultado.asignatura.nombre,
            "grade": claim.resultado.nota,
            "description": claim.descripcion,
            "status": claim.estado,
            "date": claim.fecha_solicitud,
            "deadline": claim.resultado.fecha_limite_reclamo,
            "process_year": claim.resultado.proceso.anio.year,
        } for claim in claims.order_by("-fecha_solicitud", "estudiante__apellidos")])


@extend_schema(
    tags=["Resultados"],
    summary="Resolver una reclamación de nota",
    description=(
        "Aprueba o rechaza una reclamación de nota pendiente. Solo puede resolverla el Jefe de Comisión (limitado "
        "a su provincia) o el superadmin. Al aprobar, se debe indicar fecha/hora y lugar de presentación; el "
        "estudiante es notificado."
    ),
    parameters=[
        OpenApiParameter("claim_id", OpenApiTypes.INT, OpenApiParameter.PATH, description="ID de la reclamación a resolver."),
    ],
    request=inline_serializer(
        name="ResultClaimDecisionRequest",
        fields={
            "estado": serializers.ChoiceField(choices=["aprobada", "rechazada"]),
            "fecha_presentacion": serializers.CharField(required=False, help_text="Fecha/hora ISO 8601 (requerido si estado='aprobada')."),
            "lugar_presentacion": serializers.CharField(required=False, max_length=150, help_text="Lugar de presentación (requerido si estado='aprobada')."),
        },
    ),
    responses={
        200: inline_serializer(name="ResultClaimDecisionResponse", fields={"id": serializers.IntegerField(), "status": serializers.CharField()}),
        400: inline_serializer(name="ResultClaimDecisionErrorResponse", fields={"detail": serializers.CharField()}),
        404: inline_serializer(name="ResultClaimDecisionNotFoundResponse", fields={"detail": serializers.CharField()}),
    },
    examples=[
        OpenApiExample("Reclamación resuelta", value={"success": True, "data": {"id": 5, "status": "aprobada"}, "error": None}, response_only=True),
    ],
)
class ResultClaimDecisionView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def patch(self, request, claim_id):
        if request.user.rol not in {"jefe_comision", "superadmin"}:
            return Response({"detail": "No tienes permiso para resolver reclamaciones."}, status=403)
        decision = request.data.get("estado")
        if decision not in {"aprobada", "rechazada"}:
            return Response({"detail": "El estado debe ser aprobada o rechazada."}, status=400)
        claim = Reclamacion.objects.filter(
            pk=claim_id,
            estado="pendiente",
            resultado__proceso__etapa__nombre=ETAPAS_NOMBRES[5],
        ).select_related("estudiante__escuela").first()
        if not claim:
            return Response({"detail": "La reclamación pendiente no existe."}, status=404)
        if request.user.rol == "jefe_comision" and claim.estudiante.escuela.municipio.provincia_id != request.user.provincia_id:
            return Response({"detail": "Solo puedes resolver reclamaciones de tu provincia."}, status=403)
        if decision == "aprobada":
            presentation_date = str(request.data.get("fecha_presentacion", "")).strip()
            presentation_place = str(request.data.get("lugar_presentacion", "")).strip()
            if not presentation_date or not presentation_place:
                return Response({"detail": "Para aceptar debes indicar fecha, hora y lugar de presentación."}, status=400)
            if len(presentation_place) > 150:
                return Response({"detail": "El lugar de presentación no puede superar los 150 caracteres."}, status=400)
            try:
                parsed_date = datetime.fromisoformat(presentation_date)
                if parsed_date.tzinfo is None:
                    parsed_date = timezone.make_aware(parsed_date)
            except ValueError:
                return Response({"detail": "La fecha y hora de presentación no son válidas."}, status=400)
            claim.fecha_presentacion = parsed_date
            claim.lugar_presentacion = presentation_place
        claim.estado = decision
        claim.fecha_respuesta = timezone.localdate()
        update_fields = ["estado", "fecha_respuesta"]
        if decision == "aprobada":
            update_fields.extend(["fecha_presentacion", "lugar_presentacion"])
        previous_state = claim.estado
        claim.save(update_fields=update_fields)
        record_audit(
            request.user,
            "Aprobación de reclamo de nota" if decision == "aprobada" else "Rechazo de reclamo de nota",
            request.path,
            request=request,
            previous={"estado": previous_state, "reclamacion_id": claim.id},
            new={"estado": decision, "reclamacion_id": claim.id},
        )
        if decision == "aprobada":
            from apps.core.notifications import notify_users
            presentation_text = parsed_date.strftime("%d/%m/%Y %H:%M")
            notify_users(
                [claim.estudiante.usuario],
                "Reclamación aceptada",
                f"Tu reclamación de {claim.resultado.asignatura.nombre} fue aceptada. Presentación: {presentation_text}, en {claim.lugar_presentacion}.",
            )
        return Response({"id": claim.id, "status": claim.estado})


@extend_schema(
    tags=["Resultados"],
    summary="Exportar resultados de exámenes",
    description=(
        "Exporta a Excel (.xlsx) los resultados de exámenes de ingreso para un año. Secretario/Director de "
        "escuela obtienen una hoja con Matemática/Español/Historia por estudiante de su escuela; Jefe de Comisión "
        "(acotado a su provincia) o superadmin obtienen el detalle por asignatura, filtrable por asignatura."
    ),
    parameters=[
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso. Por defecto el año actual."),
        OpenApiParameter("asignatura", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="Nombre de la asignatura para filtrar (solo aplica a Jefe de Comisión/superadmin)."),
    ],
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class ExportResultadosView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        if request.user.rol not in {"jefe_comision", "secretario_escuela", "director_escuela", "superadmin"}:
            return Response({"detail": "No tienes permiso para exportar resultados."}, status=403)
        anio = int(request.query_params.get("anio", timezone.now().year))
        if request.user.rol in {"secretario_escuela", "director_escuela"}:
            if not request.user.escuela_id:
                return Response({"detail": "El usuario no tiene una escuela asignada."}, status=400)
            escalafon_process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[1])
            if not escalafon_process:
                return Response({"detail": "No existe un escalafón para el año indicado."}, status=404)
            entries = EscalafonItem.objects.filter(
                escalafon__proceso=escalafon_process,
                escalafon__escuela_id=request.user.escuela_id,
            ).select_related("estudiante", "escalafon__escuela")
            result_process = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[5])
            results = ResultadoExamen.objects.filter(
                proceso=result_process,
                estudiante__escuela_id=request.user.escuela_id,
            ).select_related("asignatura") if result_process else ResultadoExamen.objects.none()
            results_by_student = {}
            for result in results:
                subject_key = normalize_result_subject(result.asignatura.nombre)
                if subject_key:
                    results_by_student.setdefault(result.estudiante_id, {})[subject_key] = result.nota

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Resultados"
            sheet.append(["CI", "Estudiante", "Escuela", "Matematica", "Espanol", "Historia"])
            for entry in entries.order_by("-indice_general", "estudiante__apellidos"):
                row = _school_result_row(entry, results_by_student.get(entry.estudiante_id, {}), anio)
                sheet.append([
                    row["ci"],
                    row["student"],
                    row["school"],
                    row["matematica"] if row["matematica"] is not None else "-",
                    row["espanol"] if row["espanol"] is not None else "-",
                    row["historia"] if row["historia"] is not None else "-",
                ])
            output = io.BytesIO()
            workbook.save(output)
            response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"] = f'attachment; filename="resultados_{anio}_escuela.xlsx"'
            return response

        queryset = ResultadoExamen.objects.filter(
            proceso__anio__year=anio,
            proceso__etapa__nombre=ETAPAS_NOMBRES[5],
        )
        if request.user.rol == "jefe_comision":
            queryset = queryset.filter(
                estudiante__escuela__municipio__provincia_id=request.user.provincia_id,
            )
        selected_subject = request.query_params.get("asignatura", "").strip()
        if selected_subject:
            subject_key = normalize_result_subject(selected_subject)
            subject_names = {
                "matematica": "Matemática",
                "espanol": "Español",
                "historia": "Historia",
            }
            if not subject_key:
                return Response({"detail": "La asignatura indicada no es válida."}, status=400)
            queryset = queryset.filter(asignatura__nombre__iexact=subject_names[subject_key])
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Resultados"
        sheet.append(["CI", "Asignatura", "Nota", "Fecha_Limite_Reclamo"])
        for item in queryset.select_related("estudiante", "asignatura").order_by("estudiante__apellidos", "asignatura__nombre"):
            sheet.append([item.estudiante.ci, item.asignatura.nombre, item.nota, item.fecha_limite_reclamo])
        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        filename_subject = normalize_result_subject(selected_subject) if selected_subject else "todas"
        response["Content-Disposition"] = f'attachment; filename="resultados_{filename_subject}_{anio}.xlsx"'
        return response


@extend_schema(
    tags=["Escalafón"],
    summary="Exportar el escalafón de una escuela",
    description=(
        "Exporta a Excel (.xlsx) el escalafón de una escuela para un año/proceso. Secretario/Director de escuela "
        "solo pueden exportar el de su propia escuela; Jefe de Comisión solo escuelas de su provincia."
    ),
    parameters=[
        OpenApiParameter("escuela", OpenApiTypes.STR, OpenApiParameter.QUERY, required=False, description="ID de la escuela. Por defecto la escuela del usuario autenticado."),
        OpenApiParameter("anio", OpenApiTypes.INT, OpenApiParameter.QUERY, required=False, description="Año del proceso (también acepta 'año'). Por defecto el año actual."),
    ],
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class ExportEscalafonView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        try:
            escuela = resolve_school(request.query_params.get("escuela") or request.user.escuela_id)
            if request.user.rol in {"secretario_escuela", "director_escuela"} and escuela.id != request.user.escuela_id:
                return Response({"detail": "Solo puedes exportar el escalafón de tu escuela."}, status=status.HTTP_403_FORBIDDEN)
            if request.user.rol == "jefe_comision" and escuela.municipio.provincia_id != request.user.provincia_id:
                return Response({"detail": "Solo puedes exportar escalafones de tu provincia."}, status=status.HTTP_403_FORBIDDEN)
            anio = int(request.query_params.get("anio", request.query_params.get("año", timezone.now().year)))
            proceso = Proceso.get_for_stage_and_year(anio, ETAPAS_NOMBRES[1])
            if not proceso:
                raise Escalafon.DoesNotExist
            escalafon = Escalafon.objects.get(escuela=escuela, proceso=proceso)
        except (TypeError, ValueError, Escuela.DoesNotExist, Escalafon.DoesNotExist):
            return Response({"detail": "No existe un escalafón para la escuela y año indicados."}, status=status.HTTP_404_NOT_FOUND)
        data = EscalafonExcelService().export_file(EscalafonItem.objects.filter(escalafon=escalafon).order_by("estudiante__apellidos", "estudiante__nombre"))
        response = HttpResponse(data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="escalafon-{anio}.xlsx"'
        return response


def escalafon_stage_active():
    stage = Etapa.objects.filter(nombre=ETAPAS_NOMBRES[1]).first()
    today = timezone.localdate()
    return bool(stage and stage.fecha_inicio and stage.fecha_fin and stage.fecha_inicio <= today <= stage.fecha_fin), stage


def visible_entries(request):
    process = Proceso.get_for_stage_and_year(timezone.now().year, ETAPAS_NOMBRES[1])
    entries = EscalafonItem.objects.select_related("estudiante", "escalafon__escuela").filter(escalafon__proceso=process) if process else EscalafonItem.objects.none()
    user = request.user
    if user.rol == "estudiante":
        return entries.filter(escalafon__escuela_id=user.escuela_id)
    if user.rol in {"secretario_escuela", "director_escuela"}:
        return entries.filter(escalafon__escuela_id=user.escuela_id)
    if user.rol == "jefe_comision":
        return entries.filter(escalafon__escuela__municipio__provincia_id=user.provincia_id)
    return entries


@extend_schema(
    tags=["Escalafón"],
    summary="Listar entradas del escalafón",
    description=(
        "Lista las entradas del escalafón visibles para el usuario autenticado en el proceso vigente: un "
        "estudiante ve solo su propia escuela, Secretario/Director la de su escuela, Jefe de Comisión las de su "
        "provincia, y otros roles con permiso de gestión ven todas. Las entradas se devuelven ordenadas por índice "
        "general descendente (desempate: apellidos, nombre, id) y cada una incluye `posicion` (1..n)."
    ),
    responses={
        200: inline_serializer(
            name="EscalafonListResponse",
            fields={
                "entries": EscalafonItemSerializer(many=True),  # cada entrada incluye ademas `posicion` (1..n)
                "stage_active": serializers.BooleanField(),
                "actual_id": serializers.IntegerField(allow_null=True),
            },
        ),
    },
)
class EscalafonListView(APIView):
    permission_classes = [CanManageEscalafon]

    def get(self, request):
        entries = visible_entries(request)
        active, _ = escalafon_stage_active()
        serialized = []
        for posicion, entry in rank_escalafon_entries(entries):
            item = dict(EscalafonItemSerializer(entry, context={"request": request}).data)
            item["posicion"] = posicion
            serialized.append(item)
        current = entries.filter(estudiante__usuario=request.user).first() if request.user.rol == "estudiante" else None
        return Response({"entries": serialized, "stage_active": active, "actual_id": current.id if current else None})


@extend_schema(
    tags=["Escalafón"],
    summary="Resumen provincial/escolar del escalafón",
    description=(
        "Devuelve un resumen del avance del envío del escalafón a la Comisión de Ingreso, desglosado por "
        "municipio y escuela, para el proceso del año actual. Solo accesible para Jefe de Comisión (su provincia) "
        "o Secretario de escuela (su escuela), o superusuario/superadmin."
    ),
    responses={
        200: inline_serializer(
            name="ProvincialEscalafonSummaryResponse",
            fields={
                "year": serializers.IntegerField(),
                "escuelas_enviaron": serializers.IntegerField(),
                "escuelas_pendientes": serializers.IntegerField(),
                "total_estudiantes": serializers.IntegerField(),
                "estudiantes_aceptaron": serializers.IntegerField(),
                "estudiantes_pendientes": serializers.IntegerField(),
                "municipios": serializers.ListField(child=inline_serializer(
                    name="ProvincialEscalafonMunicipio",
                    fields={
                        "id": serializers.IntegerField(),
                        "nombre": serializers.CharField(),
                        "escuelas": serializers.IntegerField(),
                        "escuelas_enviaron": serializers.IntegerField(),
                        "estudiantes": serializers.IntegerField(),
                        "estado": serializers.ChoiceField(choices=["Completo", "Parcial", "Pendiente"]),
                        "escuelas_lista": serializers.ListField(child=inline_serializer(
                            name="ProvincialEscalafonEscuela",
                            fields={
                                "id": serializers.IntegerField(),
                                "nombre": serializers.CharField(),
                                "estado": serializers.CharField(),
                            },
                        )),
                    },
                )),
            },
        ),
        403: inline_serializer(name="ProvincialEscalafonSummaryErrorResponse", fields={"detail": serializers.CharField()}),
    },
)
class ProvincialEscalafonSummaryView(APIView):
    permission_classes = [CanManageEscalafon]

    def get(self, request):
        if request.user.rol not in {"jefe_comision", "secretario_escuela"} and not request.user.is_superuser and request.user.rol != "superadmin":
            return Response({"detail": "No tienes permiso para consultar este resumen."}, status=403)
        process = Proceso.get_for_stage_and_year(timezone.now().year, ETAPAS_NOMBRES[1])
        if request.user.rol == "secretario_escuela":
            schools = Escuela.objects.filter(pk=request.user.escuela_id)
        else:
            schools = Escuela.objects.filter(municipio__provincia_id=request.user.provincia_id)
        current_entries = EscalafonItem.objects.filter(
            escalafon__proceso=process,
            escalafon__escuela__in=schools,
        ) if process else EscalafonItem.objects.none()
        municipalities = schools.values("municipio_id", "municipio__nombre").annotate(
            schools_count=Count("id"),
            sent_count=Count("escalafones", filter=Q(escalafones__proceso=process, escalafones__estado="enviado"), distinct=True),
            students_count=Count("escalafones__estudiantes", filter=Q(escalafones__proceso=process, escalafones__estado="enviado"), distinct=True),
        ).order_by("municipio__nombre")
        data = []
        for municipality in municipalities:
            sent = municipality["sent_count"]
            schools = municipality["schools_count"]
            municipality_schools = Escuela.objects.filter(
                municipio_id=municipality["municipio_id"]
            ).values("id", "nombre")
            school_statuses = dict(Escalafon.objects.filter(
                proceso=process,
                escuela_id__in=municipality_schools.values("id"),
            ).values_list("escuela_id", "estado")) if process else {}
            data.append({
                "id": municipality["municipio_id"],
                "nombre": municipality["municipio__nombre"],
                "escuelas": schools,
                "escuelas_enviaron": sent,
                "estudiantes": municipality["students_count"],
                "estado": "Completo" if sent == schools and schools else "Parcial" if sent else "Pendiente",
                "escuelas_lista": [
                    {"id": school["id"], "nombre": school["nombre"], "estado": school_statuses.get(school["id"], "pendiente")}
                    for school in municipality_schools
                ],
            })
        return Response({
            "year": timezone.now().year,
            "escuelas_enviaron": sum(item["escuelas_enviaron"] for item in data),
            "escuelas_pendientes": sum(item["escuelas"] - item["escuelas_enviaron"] for item in data),
            "total_estudiantes": current_entries.count(),
            "estudiantes_aceptaron": current_entries.filter(estado="aceptado").count(),
            "estudiantes_pendientes": current_entries.exclude(estado="aceptado").count(),
            "municipios": data,
        })


@extend_schema(
    tags=["Escalafón"],
    summary="Exportar escalafones de la provincia",
    description=(
        "Exporta a Excel (.xlsx) todas las entradas de escalafón de las escuelas de la provincia del Jefe de "
        "Comisión autenticado (o superusuario/superadmin), para el proceso del año actual."
    ),
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class ProvincialEscalafonExportView(APIView):
    permission_classes = [CanManageEscalafon]

    def get(self, request):
        if request.user.rol != "jefe_comision" and not request.user.is_superuser and request.user.rol != "superadmin":
            return Response({"detail": "Solo el Jefe de Comisión puede exportar este resumen."}, status=403)
        process = Proceso.get_for_stage_and_year(timezone.now().year, ETAPAS_NOMBRES[1])
        entries = EscalafonItem.objects.filter(
            escalafon__proceso=process,
            escalafon__escuela__municipio__provincia_id=request.user.provincia_id,
        ).order_by("escalafon__escuela__municipio__nombre", "estudiante__apellidos", "estudiante__nombre")
        data = EscalafonExcelService().export_file(entries)
        response = HttpResponse(data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="escalafones-provinciales.xlsx"'
        return response


@extend_schema(
    tags=["Escalafón"],
    summary="Descargar plantilla de escalafón",
    description="Genera una plantilla Excel (.xlsx) vacía con las columnas requeridas para importar el escalafón de una escuela.",
    responses={(200, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"): OpenApiTypes.BINARY},
)
class EscalafonTemplateView(APIView):
    permission_classes = [CanManageEscalafon]

    def get(self, request):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["CI", "Nombre", "Apellidos", "Sexo", "Dirección", "Índice_10mo", "Índice_11mo", "Índice_12mo", "Índice_General"])
        output = io.BytesIO()
        workbook.save(output)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="plantilla-escalafon.xlsx"'
        return response


@extend_schema(
    tags=["Escalafón"],
    summary="Editar una entrada del escalafón",
    description=(
        "Actualiza parcialmente (campos de índice y otros datos editables) una entrada del escalafón visible "
        "para el usuario, siempre que el escalafón de la escuela aún no haya sido enviado a la Comisión. "
        "`nombre` (máx. 150) y `apellidos` (máx. 200) se envían por separado y no pueden ser vacíos. "
        "La respuesta no incluye `posicion`; solo el listado la calcula."
    ),
    parameters=[
        OpenApiParameter("pk", OpenApiTypes.INT, OpenApiParameter.PATH, description="ID de la entrada de escalafón."),
    ],
    request=EscalafonItemSerializer,
    responses={
        200: EscalafonItemSerializer,
        403: inline_serializer(name="EscalafonEntryForbiddenResponse", fields={"detail": serializers.CharField()}),
        404: inline_serializer(name="EscalafonEntryNotFoundResponse", fields={"detail": serializers.CharField()}),
    },
)
class EscalafonEntryView(APIView):
    permission_classes = [CanManageEscalafon]

    def patch(self, request, pk):
        try:
            entry = visible_entries(request).get(pk=pk)
        except EscalafonItem.DoesNotExist:
            return Response({"detail": "Registro no encontrado."}, status=404)
        if entry.escalafon.estado == "enviado":
            return Response({"detail": "El escalafón ya fue enviado y no puede modificarse."}, status=403)
        active, _ = escalafon_stage_active()
        name_errors = {}
        for field, max_length in (("nombre", 150), ("apellidos", 200)):
            if field in request.data:
                value = request.data.get(field)
                value = value.strip() if isinstance(value, str) else ""
                if not value:
                    name_errors[field] = ["Este campo no puede estar vacío."]
                elif len(value) > max_length:
                    name_errors[field] = [f"Máximo {max_length} caracteres."]
        if name_errors:
            return Response(name_errors, status=400)
        serializer = EscalafonItemSerializer(entry, data=request.data, partial=True, context={"request": request, "stage_active": active})
        serializer.is_valid(raise_exception=True)
        previous = {field: getattr(entry, field) for field in ("indice_10", "indice_11", "indice_12", "indice_general")}
        updated = serializer.save()
        new = {field: getattr(updated, field) for field in previous}
        record_audit(request.user, "Edición de entrada del escalafón", request.path, request=request, previous=previous, new=new)
        return Response(EscalafonItemSerializer(updated, context={"request": request}).data)


@extend_schema(
    tags=["Escalafón"],
    summary="Enviar el escalafón a la Comisión",
    description=(
        "Marca como 'enviado' el escalafón de la escuela del Secretario autenticado para el proceso del año "
        "actual y bloquea la edición de los índices. Falla si existen reclamaciones ('por_revisar') pendientes. "
        "Notifica a los Jefes de Comisión de la provincia."
    ),
    request=None,
    responses={
        200: inline_serializer(name="EscalafonSendResponse", fields={"updated": serializers.IntegerField()}),
        400: inline_serializer(name="EscalafonSendErrorResponse", fields={"detail": serializers.CharField()}),
        403: inline_serializer(name="EscalafonSendForbiddenResponse", fields={"detail": serializers.CharField()}),
    },
    examples=[
        OpenApiExample("Envío exitoso", value={"success": True, "data": {"updated": 45}, "error": None}, response_only=True),
    ],
)
class EscalafonSendView(APIView):
    permission_classes = [CanManageEscalafon]

    def post(self, request):
        if request.user.rol != "secretario_escuela":
            return Response({"detail": "Solo el Secretario puede enviar índices a la Comisión."}, status=403)
        entries = visible_entries(request).filter(escalafon__proceso__anio__year=timezone.now().year)
        if entries.filter(estado="por_revisar").exists():
            return Response({"detail": "No puedes enviar el escalafón mientras existan reclamaciones pendientes."}, status=400)
        escalafones = Escalafon.objects.filter(
            id__in=entries.values("escalafon_id"),
            escuela=request.user.escuela,
        )
        with transaction.atomic():
            escalafones.update(estado="enviado")
            entries.update(indices_bloqueados=True)
        record_audit(
            request.user,
            "Envío de escalafón a Comisión",
            request.path,
            request=request,
            new={"escuela": request.user.escuela.nombre, "entradas": entries.count()},
        )
        from apps.core.notifications import notify_users
        notify_users(
            Usuario.objects.filter(rol="jefe_comision", provincia_id=request.user.provincia_id),
            "Escalafón enviado a comisión",
            f"La escuela {request.user.escuela.nombre} envió su escalafón a la Comisión de Ingreso.",
        )
        return Response({"updated": entries.count()})


@extend_schema(
    tags=["Escalafón"],
    summary="Aceptar o solicitar revisión del propio escalafón",
    description=(
        "Permite a un estudiante autenticado aceptar sus índices de escalafón del proceso del año actual "
        "(`action=aceptar`) o solicitar una revisión indicando una causa (`action=revision`), siempre que el "
        "escalafón de su escuela aún no haya sido enviado a la Comisión. Notifica al Secretario de la escuela."
    ),
    parameters=[
        OpenApiParameter("action", OpenApiTypes.STR, OpenApiParameter.PATH, enum=["aceptar", "revision"], description="Acción a realizar."),
    ],
    request=StudentEscalafonActionSerializer,
    responses={
        200: EscalafonItemSerializer,
        400: inline_serializer(name="StudentEscalafonActionErrorResponse", fields={"detail": serializers.CharField()}),
        403: inline_serializer(name="StudentEscalafonActionForbiddenResponse", fields={"detail": serializers.CharField()}),
        404: inline_serializer(name="StudentEscalafonActionNotFoundResponse", fields={"detail": serializers.CharField()}),
    },
)
class StudentEscalafonActionView(APIView):
    permission_classes = [CanManageEscalafon]

    def post(self, request, action):
        if request.user.rol != "estudiante":
            return Response({"detail": "Solo un estudiante puede realizar esta acción."}, status=403)
        try:
            entry = EscalafonItem.objects.get(estudiante__usuario=request.user, escalafon__proceso__anio__year=timezone.now().year)
        except EscalafonItem.DoesNotExist:
            return Response({"detail": "No tienes un escalafón vigente."}, status=404)
        serializer = StudentEscalafonActionSerializer(data=request.data, context={"action": action})
        serializer.is_valid(raise_exception=True)
        if action not in {"aceptar", "revision"}:
            return Response({"detail": "Acción no válida."}, status=400)
        if entry.escalafon.estado == "enviado":
            return Response({"detail": "El escalafón ya fue enviado y solo puede consultarse."}, status=403)
        previous_state = entry.estado
        entry.estado = "aceptado" if action == "aceptar" else "por_revisar"
        entry.causa_revision = serializer.validated_data.get("causa", "") if action == "revision" else ""
        entry.fecha_revision = timezone.now() if action == "revision" else None
        entry.save(update_fields=["estado", "causa_revision", "fecha_revision"])
        record_audit(
            request.user,
            "Aceptación de escalafón" if action == "aceptar" else "Reclamo de escalafón",
            request.path,
            request=request,
            previous={"estado": previous_state},
            new={"estado": entry.estado, "causa": entry.causa_revision},
        )
        from apps.core.notifications import notify_users
        secretaries = Usuario.objects.filter(rol="secretario_escuela", escuela=entry.escalafon.escuela)
        message = f"El estudiante {entry.estudiante.nombre} {entry.estudiante.apellidos} {'solicitó revisión' if action == 'revision' else 'aceptó sus índices'}."
        notify_users(secretaries, "Solicitud de revisión", message)
        return Response(EscalafonItemSerializer(entry, context={"request": request}).data)


@extend_schema(
    tags=["Escalafón"],
    summary="Atender una solicitud de revisión del escalafón",
    description=(
        "Marca como atendida ('sin_respuesta') una entrada de escalafón que estaba en estado 'por_revisar'. Solo "
        "puede ejecutarlo el Secretario de la escuela correspondiente. Notifica al estudiante."
    ),
    parameters=[
        OpenApiParameter("pk", OpenApiTypes.INT, OpenApiParameter.PATH, description="ID de la entrada de escalafón."),
    ],
    request=None,
    responses={
        200: EscalafonItemSerializer,
        400: inline_serializer(name="EscalafonReviewErrorResponse", fields={"detail": serializers.CharField()}),
        403: inline_serializer(name="EscalafonReviewForbiddenResponse", fields={"detail": serializers.CharField()}),
        404: inline_serializer(name="EscalafonReviewNotFoundResponse", fields={"detail": serializers.CharField()}),
    },
)
class EscalafonReviewView(APIView):
    permission_classes = [CanManageEscalafon]

    def post(self, request, pk):
        if request.user.rol != "secretario_escuela":
            return Response({"detail": "Solo el Secretario puede revisar solicitudes."}, status=403)
        try:
            entry = visible_entries(request).get(pk=pk)
        except EscalafonItem.DoesNotExist:
            return Response({"detail": "Registro no encontrado."}, status=404)
        if entry.estado != "por_revisar":
            return Response({"detail": "Esta reclamación no está pendiente."}, status=400)
        previous_state = entry.estado
        entry.estado = "sin_respuesta"
        entry.save(update_fields=["estado"])
        record_audit(
            request.user,
            "Atención de reclamo de escalafón",
            request.path,
            request=request,
            previous={"estado": previous_state},
            new={"estado": entry.estado, "entrada_id": entry.id},
        )
        from apps.core.notifications import notify_users
        notify_users(
            [entry.estudiante.usuario],
            "Revisión atendida",
            "El secretario revisó tu solicitud de revisión del escalafón.",
        )
        return Response(EscalafonItemSerializer(entry, context={"request": request}).data)


@extend_schema(
    tags=["Importación y exportación"],
    summary="Consultar el estado de una importación",
    description=(
        "Las importaciones Excel se ejecutan de forma asíncrona (Celery). Este endpoint devuelve el estado de la "
        "tarea encolada y, al terminar, el resultado con el mismo cuerpo y código HTTP que antes devolvía la "
        "importación síncrona. Solo el usuario que inició la importación puede consultarla (404 en otro caso). "
        "El identificador caduca a la hora."
    ),
    parameters=[
        OpenApiParameter("task_id", OpenApiTypes.STR, OpenApiParameter.PATH, description="Identificador devuelto por la importación (202)."),
    ],
    responses={
        200: inline_serializer(
            name="ImportTaskStatusResponse",
            fields={
                "estado": serializers.ChoiceField(choices=["pendiente", "en_proceso", "completada", "fallida"]),
                "http_status": serializers.IntegerField(allow_null=True),
                "resultado": serializers.DictField(allow_null=True),
            },
        ),
        404: inline_serializer(name="ImportTaskNotFoundResponse", fields={"detail": serializers.CharField()}),
    },
    examples=[
        OpenApiExample(
            "Importación completada",
            value={"success": True, "data": {"estado": "completada", "http_status": 201, "resultado": {"inserted": 12, "updated": 3, "errors": []}}, "error": None},
            response_only=True,
        ),
        OpenApiExample(
            "Importación en curso",
            value={"success": True, "data": {"estado": "en_proceso", "http_status": None, "resultado": None}, "error": None},
            response_only=True,
        ),
    ],
)
class ImportTaskStatusView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, task_id):
        from celery.result import AsyncResult
        from django.core.cache import cache

        meta = cache.get(f"import_task:{task_id}")
        if not meta or meta.get("user_id") != request.user.id:
            return Response({"detail": "Tarea no encontrada."}, status=status.HTTP_404_NOT_FOUND)
        payload = cache.get(f"import_result:{task_id}")
        if payload is None:
            async_result = AsyncResult(str(task_id))
            state = async_result.state
            if state == "SUCCESS":
                payload = async_result.result
            elif state in {"FAILURE", "REVOKED"}:
                payload = {"http_status": 500, "body": {"detail": "No se pudo procesar el archivo Excel. Inténtelo de nuevo."}}
            else:
                estado = "en_proceso" if state in {"STARTED", "RETRY"} else "pendiente"
                return Response({"estado": estado, "http_status": None, "resultado": None})
        return Response({
            "estado": "completada" if payload["http_status"] < 400 else "fallida",
            "http_status": payload["http_status"],
            "resultado": payload["body"],
        })
