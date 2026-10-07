from datetime import date, timedelta
import tempfile
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APITestCase

from apps.authentication.models import Estudiante, Usuario
from apps.gestion_escuela.models import Escalafon, EscalafonItem
from apps.gestion_personal.models import BoletaInteres, BoletaInteresItem, BoletaSolicitud, BoletaSolicitudItem
from apps.superadmin.models import Carrera, Ces, Escuela, Municipio, Provincia, TipoOtorgamiento

from .models import ETAPAS_NOMBRES, Etapa, Noticia, PlanPlaza, Proceso


class ProvincialNewsTests(APITestCase):
    @classmethod
    def setUpClass(cls):
        cls.media_directory = tempfile.TemporaryDirectory()
        cls.media_settings = override_settings(MEDIA_ROOT=cls.media_directory.name)
        cls.media_settings.enable()
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        cls.media_settings.disable()
        cls.media_directory.cleanup()

    def setUp(self):
        self.holguin = Provincia.objects.create(nombre="Holguín")
        self.matanzas = Provincia.objects.create(nombre="Matanzas")
        self.manager = Usuario.objects.create_user(
            username="news.manager",
            email="news.manager@example.com",
            password="secret1234",
            rol="gestor_noticias",
            provincia=self.matanzas,
        )

    def create_news(self, province, **fields):
        values = {
            "titulo": "Aviso provincial",
            "contenido": "Información importante.",
            "permanente": False,
            "fecha_expiracion": timezone.localdate() + timedelta(days=2),
        }
        values.update(fields)
        return Noticia.objects.create(provincia=province, **values)

    def test_public_landing_defaults_to_holguin_and_excludes_expired_news(self):
        active = self.create_news(self.holguin)
        permanent = self.create_news(self.holguin, titulo="Aviso permanente", permanente=True, fecha_expiracion=None)
        self.create_news(self.holguin, titulo="Aviso vencido", fecha_expiracion=timezone.localdate() - timedelta(days=1))
        self.create_news(self.matanzas, titulo="Otra provincia")

        response = self.client.get(reverse("noticias-landing"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["provincia_nombre"], "Holguín")
        self.assertEqual({item["id"] for item in response.data["items"]}, {active.id, permanent.id})
        other_province_news = Noticia.objects.get(titulo="Otra provincia")
        download = self.client.get(reverse("noticia-multimedia-download", args=[other_province_news.id]))
        self.assertEqual(download.status_code, 404)

    def test_authenticated_landing_uses_the_users_province(self):
        news = self.create_news(self.matanzas)
        self.client.force_authenticate(self.manager)

        response = self.client.get(reverse("noticias-landing"))

        self.assertEqual(response.data["provincia_nombre"], "Matanzas")
        self.assertEqual([item["id"] for item in response.data["items"]], [news.id])

    def test_manager_can_download_expired_news_but_public_users_cannot(self):
        expired = self.create_news(
            self.matanzas,
            fecha_expiracion=timezone.localdate() - timedelta(days=1),
            multimedia=SimpleUploadedFile("aviso.pdf", b"%PDF-1.4 aviso"),
        )
        self.client.force_authenticate(self.manager)
        manager_response = self.client.get(reverse("noticia-multimedia-download", args=[expired.id]))
        self.assertEqual(manager_response.status_code, 200)

        self.client.force_authenticate(None)
        public_response = self.client.get(reverse("noticia-multimedia-download", args=[expired.id]))
        self.assertEqual(public_response.status_code, 404)

    def test_manager_crud_is_limited_to_assigned_province(self):
        other_news = self.create_news(self.holguin)
        self.client.force_authenticate(self.manager)
        created = self.client.post(reverse("provincial-noticias"), {
            "titulo": "Comunicado de Matanzas",
            "contenido": "Contenido de prueba.",
            "fecha_expiracion": (timezone.localdate() + timedelta(days=4)).isoformat(),
            "permanente": "false",
        }, format="multipart")

        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["provincia"], self.matanzas.id)
        listed = self.client.get(reverse("provincial-noticias"))
        self.assertEqual([item["id"] for item in listed.data], [created.data["id"]])
        self.assertNotEqual(created.data["id"], other_news.id)

        updated = self.client.patch(
            reverse("provincial-noticia-detail", args=[created.data["id"]]),
            {"titulo": "Comunicado actualizado", "permanente": "true"},
            format="multipart",
        )
        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertTrue(updated.data["permanente"])
        self.assertIsNone(updated.data["fecha_expiracion"])

        deleted = self.client.delete(reverse("provincial-noticia-detail", args=[created.data["id"]]))
        self.assertEqual(deleted.status_code, 204)

    def test_non_news_manager_cannot_use_management_api(self):
        student = Usuario.objects.create_user(
            username="news.student",
            email="news.student@example.com",
            password="secret1234",
            rol="estudiante",
            provincia=self.matanzas,
        )
        self.client.force_authenticate(student)

        response = self.client.get(reverse("provincial-noticias"))

        self.assertEqual(response.status_code, 403)

    def test_commission_chief_can_manage_news_only_in_assigned_province(self):
        chief = Usuario.objects.create_user(
            username="news.chief",
            email="news.chief@example.com",
            password="secret1234",
            rol="jefe_comision",
            provincia=self.matanzas,
        )
        holguin_news = self.create_news(self.holguin, titulo="Aviso de Holguín")
        expired_news = self.create_news(
            self.matanzas,
            titulo="Aviso vencido del jefe",
            fecha_expiracion=timezone.localdate() - timedelta(days=1),
            multimedia=SimpleUploadedFile("aviso-vencido.pdf", b"%PDF-1.4 aviso"),
        )
        self.client.force_authenticate(chief)
        expired_download = self.client.get(
            reverse("noticia-multimedia-download", args=[expired_news.id])
        )
        self.assertEqual(expired_download.status_code, 200)
        self.assertEqual(b"".join(expired_download.streaming_content), b"%PDF-1.4 aviso")

        created = self.client.post(reverse("provincial-noticias"), {
            "titulo": "Aviso del jefe de comisión",
            "contenido": "El jefe también puede administrar noticias de su provincia.",
            "permanente": "true",
        }, format="multipart")

        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["provincia"], self.matanzas.id)
        listed = self.client.get(reverse("provincial-noticias"))
        self.assertEqual({item["id"] for item in listed.data}, {created.data["id"], expired_news.id})
        self.assertNotIn(holguin_news.id, [item["id"] for item in listed.data])

        deleted = self.client.delete(reverse("provincial-noticia-detail", args=[created.data["id"]]))
        self.assertEqual(deleted.status_code, 204)

    def test_uploaded_document_can_be_downloaded_and_invalid_types_are_rejected(self):
        self.client.force_authenticate(self.manager)
        created = self.client.post(reverse("provincial-noticias"), {
            "titulo": "Documento provincial",
            "contenido": "Descarga el documento adjunto.",
            "permanente": "true",
            "multimedia": SimpleUploadedFile(
                "orientaciones.pdf", b"%PDF-1.4 contenido", content_type="application/pdf"
            ),
        }, format="multipart")

        self.assertEqual(created.status_code, 201, created.data)
        download = self.client.get(reverse("noticia-multimedia-download", args=[created.data["id"]]))
        self.assertEqual(download.status_code, 200)
        self.assertIn("attachment", download["Content-Disposition"])
        self.assertEqual(b"".join(download.streaming_content), b"%PDF-1.4 contenido")

        invalid = self.client.post(reverse("provincial-noticias"), {
            "titulo": "Archivo inválido",
            "contenido": "No debe aceptar ejecutables.",
            "permanente": "true",
            "multimedia": SimpleUploadedFile(
                "programa.exe", b"MZ executable", content_type="application/octet-stream"
            ),
        }, format="multipart")
        self.assertEqual(invalid.status_code, 400)
        self.assertIn("multimedia", invalid.data)

        invalid_signature = self.client.post(reverse("provincial-noticias"), {
            "titulo": "PDF falso",
            "contenido": "El contenido debe coincidir con la extensión.",
            "permanente": "true",
            "multimedia": SimpleUploadedFile(
                "falso.pdf", b"MZ executable", content_type="application/pdf"
            ),
        }, format="multipart")
        self.assertEqual(invalid_signature.status_code, 400)

    def test_uploaded_video_has_a_safe_inline_preview_and_download(self):
        self.client.force_authenticate(self.manager)
        created = self.client.post(reverse("provincial-noticias"), {
            "titulo": "Video provincial",
            "contenido": "Video del proceso.",
            "permanente": "true",
            "multimedia": SimpleUploadedFile(
                "proceso.mp4", b"\x00\x00\x00\x18ftypisom0000", content_type="video/mp4"
            ),
        }, format="multipart")

        self.assertEqual(created.status_code, 201, created.data)
        landing = self.client.get(reverse("noticias-landing"))
        item = landing.data["items"][0]
        self.assertEqual(item["media_type"], "video")
        self.assertIn("inline=1", item["media_url"])
        self.assertTrue(item["media_url"].startswith("http://testserver/"))
        inline = self.client.get(
            reverse("noticia-multimedia-download", args=[created.data["id"]]),
            {"inline": "1"},
        )
        self.assertEqual(inline.status_code, 200)
        self.assertEqual(inline["Content-Type"], "video/mp4")
        self.assertEqual(b"".join(inline.streaming_content), b"\x00\x00\x00\x18ftypisom0000")

    @patch("apps.gestion_provincial.serializers_noticias.fetch_open_graph_image")
    def test_external_article_cover_is_cached_for_landing(self, fetch_cover):
        fetch_cover.return_value = (b"\x89PNG\r\n\x1a\npreview", ".png")
        self.client.force_authenticate(self.manager)
        created = self.client.post(reverse("provincial-noticias"), {
            "titulo": "Enlace externo",
            "contenido": "Lee el comunicado completo.",
            "url_externa": "https://example.org/noticia",
            "permanente": "true",
            "multimedia": SimpleUploadedFile(
                "comunicado.pdf", b"%PDF-1.4 contenido", content_type="application/pdf"
            ),
        }, format="multipart")

        self.assertEqual(created.status_code, 201, created.data)
        self.assertIn("tipo=portada", created.data["imagen_portada_url"])
        landing = self.client.get(reverse("noticias-landing"))
        self.assertIn("tipo=portada", landing.data["items"][0]["cover_url"])
        cover = self.client.get(
            reverse("noticia-multimedia-download", args=[created.data["id"]]),
            {"tipo": "portada", "inline": "1"},
        )
        self.assertEqual(cover.status_code, 200)
        self.assertEqual(cover["Content-Type"], "image/png")
        self.assertEqual(cover["X-Content-Type-Options"], "nosniff")
        fetch_cover.assert_called_once_with("https://example.org/noticia")

    @patch("apps.gestion_provincial.services.noticias._safe_get")
    def test_open_graph_image_is_extracted_and_signature_checked(self, safe_get):
        html = b'<meta property="og:image" content="/media/cover.png">'
        png = b"\x89PNG\r\n\x1a\nvalid-preview"
        safe_get.side_effect = [
            ("https://example.org/noticia", "text/html", html),
            ("https://example.org/media/cover.png", "image/png", png),
        ]

        from apps.gestion_provincial.services.noticias import fetch_open_graph_image

        self.assertEqual(fetch_open_graph_image("https://example.org/noticia"), (png, ".png"))

    def test_open_graph_fetch_rejects_private_ip_addresses(self):
        from apps.gestion_provincial.services.noticias import fetch_open_graph_image

        self.assertIsNone(fetch_open_graph_image("http://127.0.0.1/internal"))


class EtapaActivationTests(APITestCase):
    def setUp(self):
        self.user = Usuario.objects.create_user(
            username="jefe_comision",
            email="jefe@example.com",
            password="password123",
            rol="jefe_comision",
        )
        self.client.force_authenticate(self.user)
        self.stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        self.proceso = Proceso.objects.create(
            anio=date(2025, 1, 1),
            etapa=self.stage,
        )

    def test_creates_six_stages_for_a_process(self):
        self.assertEqual(Etapa.objects.count(), 6)
        self.assertEqual(Etapa.objects.values_list("nombre", flat=True).count(), 6)
        self.assertEqual(set(Etapa.objects.values_list("nombre", flat=True)), set(ETAPAS_NOMBRES.values()))

    def test_rejects_stage_names_outside_catalog(self):
        with self.assertRaises(Exception):
            Etapa.objects.create(nombre="Etapa inventada")

    def test_cannot_activate_stage_out_of_sequence(self):
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[2])

        response = self.client.post(
            reverse("provincial-etapa-activar", args=[stage.pk]),
            {"fecha_inicio": "2025-01-02", "fecha_fin": "2025-01-10"},
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("secuencia", response.data["detail"])

    def test_activation_returns_current_status(self):
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        today = date.today()
        response = self.client.post(
            reverse("provincial-etapa-activar", args=[stage.pk]),
            {
                "fecha_inicio": today.isoformat(),
                "fecha_fin": (today + timedelta(days=5)).isoformat(),
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["estado"], "en_curso")

    def test_activation_does_not_require_a_process(self):
        self.proceso.delete()
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        today = date.today()

        response = self.client.post(
            reverse("provincial-etapa-activar", args=[stage.pk]),
            {
                "fecha_inicio": today.isoformat(),
                "fecha_fin": (today + timedelta(days=5)).isoformat(),
            },
            format="json",
        )

        self.assertEqual(response.status_code, 200)

    def test_can_close_stage_before_planned_end_date(self):
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        today = date.today()
        self.client.post(
            reverse("provincial-etapa-activar", args=[stage.pk]),
            {"fecha_inicio": today.isoformat(), "fecha_fin": (today + timedelta(days=30)).isoformat()},
            format="json",
        )

        response = self.client.post(reverse("provincial-etapa-cerrar", args=[stage.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["estado"], "completada")

    def test_can_close_stage_before_its_start_date(self):
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        today = date.today()
        planned_start = today + timedelta(days=1)
        planned_end = today + timedelta(days=10)
        self.client.post(
            reverse("provincial-etapa-activar", args=[stage.pk]),
            {"fecha_inicio": planned_start.isoformat(), "fecha_fin": planned_end.isoformat()},
            format="json",
        )

        response = self.client.post(reverse("provincial-etapa-cerrar", args=[stage.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["estado"], "completada")
        self.assertEqual(response.data["fecha_fin"], planned_end.isoformat())

    def test_closing_stage_auto_sends_pending_escalafones(self):
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        today = date.today()
        self.client.post(
            reverse("provincial-etapa-activar", args=[stage.pk]),
            {"fecha_inicio": today.isoformat(), "fecha_fin": (today + timedelta(days=10)).isoformat()},
            format="json",
        )

        provincia = Provincia.objects.create(nombre="Provincia cerrada")
        municipio = Municipio.objects.create(nombre="Municipio cerrado", provincia=provincia)
        escuela_pending = Escuela.objects.create(nombre="Escuela pendiente", municipio=municipio)
        escuela_sent = Escuela.objects.create(nombre="Escuela enviada", municipio=municipio)

        escalafon_pending = Escalafon.objects.create(proceso=self.proceso, escuela=escuela_pending, estado="pendiente")
        escalafon_sent = Escalafon.objects.create(proceso=self.proceso, escuela=escuela_sent, estado="enviado")

        estudiante_pending = Estudiante.objects.create(
            ci="10000000001", nombre="Ana", apellidos="Pendiente", sexo="F",
            direccion="Calle 1", escuela=escuela_pending,
        )
        estudiante_sent = Estudiante.objects.create(
            ci="10000000002", nombre="Luis", apellidos="Enviado", sexo="M",
            direccion="Calle 2", escuela=escuela_sent,
        )

        item_pending = EscalafonItem.objects.create(
            escalafon=escalafon_pending, estudiante=estudiante_pending,
            indice_10=80, indice_11=82, indice_12=84, indice_general=82,
        )
        item_sent = EscalafonItem.objects.create(
            escalafon=escalafon_sent, estudiante=estudiante_sent,
            indice_10=90, indice_11=92, indice_12=93, indice_general=92,
        )

        response = self.client.post(reverse("provincial-etapa-cerrar", args=[stage.pk]))

        self.assertEqual(response.status_code, 200)
        escalafon_pending.refresh_from_db()
        escalafon_sent.refresh_from_db()
        item_pending.refresh_from_db()
        item_sent.refresh_from_db()
        self.assertEqual(escalafon_pending.estado, "enviado")
        self.assertEqual(escalafon_sent.estado, "enviado")
        self.assertTrue(item_pending.indices_bloqueados)
        self.assertFalse(item_sent.indices_bloqueados)

    def test_rejects_end_date_before_start_date(self):
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        response = self.client.post(
            reverse("provincial-etapa-activar", args=[stage.pk]),
            {"fecha_inicio": "2026-02-10", "fecha_fin": "2026-02-01"},
            format="json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("fecha", str(response.data).lower())

    def test_expired_stage_is_closed_when_stages_are_listed(self):
        stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        stage.fecha_inicio = date.today() - timedelta(days=5)
        stage.fecha_fin = date.today() - timedelta(days=1)
        stage.estado = "en_curso"
        stage.save()

        response = self.client.get(reverse("provincial-etapas"))

        self.assertEqual(response.status_code, 200)
        stage.refresh_from_db()
        self.assertEqual(stage.estado, "completada")

    def test_cannot_reset_until_all_stages_are_completed(self):
        response = self.client.post(reverse("provincial-etapas-reiniciar"))

        self.assertEqual(response.status_code, 400)

    def test_reset_clears_dates_and_states_after_all_stages_complete(self):
        Etapa.objects.update(
            fecha_inicio=date.today(), fecha_fin=date.today(), estado="completada"
        )

        response = self.client.post(reverse("provincial-etapas-reiniciar"))

        self.assertEqual(response.status_code, 200)
        self.assertFalse(Etapa.objects.exclude(fecha_inicio=None).exists())
        self.assertFalse(Etapa.objects.exclude(fecha_fin=None).exists())
        self.assertFalse(Etapa.objects.exclude(estado="no_iniciada").exists())

    def test_provincial_representative_cannot_manage_stages(self):
        provincial = Usuario.objects.create_user(
            username="provincial",
            email="provincial@example.com",
            password="password123",
            rol="ingreso_provincial",
        )
        self.client.force_authenticate(provincial)
        response = self.client.get(reverse("provincial-etapas"))

        self.assertEqual(response.status_code, 403)

    def test_secretary_can_read_stage_status(self):
        secretary = Usuario.objects.create_user(
            username="secretario",
            email="secretario@example.com",
            password="password123",
            rol="secretario_escuela",
        )
        self.client.force_authenticate(secretary)
        response = self.client.get(reverse("provincial-etapas"))

        self.assertEqual(response.status_code, 200)

    def test_student_can_read_stage_status(self):
        student = Usuario.objects.create_user(
            username="estudiante",
            email="estudiante@example.com",
            password="password123",
            rol="estudiante",
        )
        self.client.force_authenticate(student)
        response = self.client.get(reverse("provincial-etapas"))

        self.assertEqual(response.status_code, 200)

    def test_commission_chief_can_read_real_dashboard_metrics(self):
        response = self.client.get(reverse("provincial-dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["municipios"], 0)
        self.assertEqual(response.data["estudiantes"], 0)
        self.assertIn("etapa_activa", response.data)
        self.assertIn("top_carreras", response.data)

    def test_commission_chief_dashboard_is_limited_to_own_province(self):
        own_province = Provincia.objects.create(nombre="Provincia propia")
        other_province = Provincia.objects.create(nombre="Otra provincia")
        own_municipality = Municipio.objects.create(nombre="Municipio propio", provincia=own_province)
        other_municipality = Municipio.objects.create(nombre="Otro municipio", provincia=other_province)
        Escuela.objects.create(nombre="Escuela propia", municipio=own_municipality)
        Escuela.objects.create(nombre="Escuela ajena", municipio=other_municipality)
        self.user.provincia = own_province
        self.user.save(update_fields=["provincia"])

        response = self.client.get(reverse("provincial-dashboard"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["municipios"], 1)
        self.assertEqual(response.data["escuelas"], 1)
        self.assertEqual(response.data["municipios_lista"][0]["nombre"], "Municipio propio")

    def test_dashboard_top_careers_uses_active_stage_and_own_province(self):
        own_province = Provincia.objects.create(nombre="Provincia del jefe")
        other_province = Provincia.objects.create(nombre="Otra provincia")
        own_school = Escuela.objects.create(
            nombre="Escuela propia",
            municipio=Municipio.objects.create(nombre="Municipio propio", provincia=own_province),
        )
        other_school = Escuela.objects.create(
            nombre="Escuela ajena",
            municipio=Municipio.objects.create(nombre="Municipio ajeno", provincia=other_province),
        )
        self.user.provincia = own_province
        self.user.save(update_fields=["provincia"])

        own_student = Estudiante.objects.create(
            ci="10000000001", nombre="Estudiante", apellidos="Propio", sexo="F",
            direccion="Calle 1", escuela=own_school,
        )
        other_student = Estudiante.objects.create(
            ci="10000000002", nombre="Estudiante", apellidos="Ajeno", sexo="M",
            direccion="Calle 2", escuela=other_school,
        )
        ces = Ces.objects.create(nombre="CES de prueba")
        interest_career = Carrera.objects.create(
            codigo="INT-1", nombre="Carrera de interés", ces=ces, provincia=own_province,
        )
        request_career = Carrera.objects.create(
            codigo="SOL-1", nombre="Carrera de solicitud", ces=ces, provincia=own_province,
        )
        interest_process = Proceso.objects.create(
            anio=date(timezone.now().year, 1, 1), etapa=Etapa.objects.get(nombre=ETAPAS_NOMBRES[2]),
        )
        own_interest = BoletaInteres.objects.create(
            estudiante=own_student, proceso=interest_process, enviada=True,
        )
        other_interest = BoletaInteres.objects.create(
            estudiante=other_student, proceso=interest_process, enviada=True,
        )
        BoletaInteresItem.objects.create(boleta_interes=own_interest, carrera=interest_career, prioridad=1)
        BoletaInteresItem.objects.create(boleta_interes=other_interest, carrera=request_career, prioridad=1)

        stage_2 = Etapa.objects.get(nombre=ETAPAS_NOMBRES[2])
        stage_2.estado = "en_curso"
        stage_2.save(update_fields=["estado"])
        interest_response = self.client.get(reverse("provincial-dashboard"))

        self.assertEqual(interest_response.status_code, 200)
        self.assertEqual(interest_response.data["top_carreras"], [{"carrera__nombre": "Carrera de interés", "total": 1}])

        stage_2.estado = "completada"
        stage_2.save(update_fields=["estado"])
        stage_3 = Etapa.objects.get(nombre=ETAPAS_NOMBRES[3])
        stage_3.estado = "en_curso"
        stage_3.save(update_fields=["estado"])
        request_process = Proceso.objects.create(
            anio=date(timezone.now().year, 1, 1), etapa=stage_3,
        )
        tipo_provincial, _ = TipoOtorgamiento.objects.get_or_create(nombre="Provincial")
        plan = PlanPlaza.objects.create(
            proceso=request_process, carrera=request_career, cantidad_plazas=1,
            otorgamiento_tipo=tipo_provincial, ces=ces, provincia=own_province, sexo="A",
        )
        request_ballot = BoletaSolicitud.objects.create(
            estudiante=own_student, proceso=request_process, estado="pendiente",
        )
        BoletaSolicitudItem.objects.create(boleta_solicitud=request_ballot, plan_plaza=plan, prioridad=1)
        other_plan = PlanPlaza.objects.create(
            proceso=request_process, carrera=interest_career, cantidad_plazas=1,
            otorgamiento_tipo=tipo_provincial, ces=ces, provincia=other_province, sexo="A",
        )
        other_request = BoletaSolicitud.objects.create(
            estudiante=other_student, proceso=request_process, estado="aprobada",
        )
        BoletaSolicitudItem.objects.create(boleta_solicitud=other_request, plan_plaza=other_plan, prioridad=1)

        request_response = self.client.get(reverse("provincial-dashboard"))

        self.assertEqual(request_response.status_code, 200)
        self.assertEqual(request_response.data["top_carreras"], [{"carrera__nombre": "Carrera de solicitud", "total": 1}])

    def test_director_can_read_stage_status(self):
        director = Usuario.objects.create_user(
            username="director",
            email="director@example.com",
            password="password123",
            rol="director_escuela",
        )
        self.client.force_authenticate(director)
        response = self.client.get(reverse("provincial-etapas"))

        self.assertEqual(response.status_code, 200)

    def test_landing_can_read_stage_availability_without_authentication(self):
        self.client.force_authenticate(user=None)

        response = self.client.get(reverse("etapas-disponibilidad"))

        self.assertEqual(response.status_code, 200)
        self.assertIn("etapas", response.data)
        self.assertIn("registro_estudiantil", response.data)


class PlanPlazaImportTests(APITestCase):
    def setUp(self):
        self.user = Usuario.objects.create_user(
            username="jefe",
            email="jefe@example.com",
            password="password123",
            rol="jefe_comision",
        )
        self.client.force_authenticate(self.user)
        self.provincia = Provincia.objects.create(nombre="La Habana")
        self.ces = self._create_ces("Universidad de La Habana")
        self.carrera = self._create_carrera("CS-101", "Ingeniería en Informática")
        self.active_stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[1])
        self.active_stage.estado = "en_curso"
        self.active_stage.save(update_fields=["estado"])
        self.proceso = Proceso.objects.create(anio=date(2025, 1, 1), etapa=self.active_stage)
        Etapa.objects.filter(nombre=ETAPAS_NOMBRES[3]).update(estado="en_curso")

    def _create_ces(self, nombre):
        from apps.superadmin.models import Ces
        return Ces.objects.create(nombre=nombre, activa=True)

    def _create_carrera(self, codigo, nombre):
        from apps.superadmin.models import Carrera
        return Carrera.objects.create(
            codigo=codigo,
            nombre=nombre,
            ces=self.ces,
            provincia=self.provincia,
            activa=True,
        )

    def _import_plan_plaza(self, file):
        """Encola la importación (202) y devuelve la respuesta de /tareas/<id>/ ya resuelta (modo eager)."""
        accepted = self.client.post(reverse("import-plan-plaza"), {"file": file}, format="multipart")
        self.assertEqual(accepted.status_code, 202, accepted.content)
        task_id = accepted.data["task_id"]
        status_response = self.client.get(reverse("import-task-status", args=[task_id]))
        self.assertEqual(status_response.status_code, 200)
        data = status_response.data
        self.assertIn(data["estado"], {"completada", "fallida"})
        response = type("R", (), {})()
        response.status_code = data["http_status"]
        response.data = data["resultado"]
        response.content = str(data).encode()
        return response

    def test_imports_plan_plaza_using_career_pk_and_active_process(self):
        from io import BytesIO
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plan de Plazas"
        sheet.append(["Código_Carrera", "Nombre_Carrera", "Cantidad_Plazas", "Tipo_Otorgamiento", "CES", "Provincia", "Sexo"])
        sheet.append([self.carrera.codigo, self.carrera.nombre, 12, "Municipal", self.ces.nombre, self.provincia.nombre, "A"])

        file = SimpleUploadedFile(
            "plan-plazas.xlsx",
            BytesIO().getvalue() if False else b"",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        buffer = BytesIO()
        workbook.save(buffer)
        file = SimpleUploadedFile("plan-plazas.xlsx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self._import_plan_plaza(file)

        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.data["inserted"], 1)

        plan = self.carrera.plan_plaza.order_by("-id").first()
        self.assertIsNotNone(plan)
        self.assertEqual(plan.carrera_id, self.carrera.id)
        self.assertEqual(plan.otorgamiento_tipo.nombre, "Municipal")
        plan_stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[3])
        self.assertEqual(plan.proceso.etapa_id, plan_stage.id)
        self.assertEqual(plan.cantidad_plazas, 12)
        self.assertEqual(plan.sexo, "A")

    def test_import_uses_user_province_and_creates_process_automatically(self):
        from io import BytesIO
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.user.provincia = self.provincia
        self.user.save(update_fields=["provincia"])

        plan_stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[3])
        plan_stage.estado = "en_curso"
        plan_stage.save(update_fields=["estado"])
        Proceso.objects.filter(etapa=plan_stage).delete()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plan de Plazas"
        sheet.append(["Código_Carrera", "Nombre_Carrera", "Cantidad_Plazas", "Tipo_Otorgamiento", "CES", "Provincia", "Sexo"])
        sheet.append([self.carrera.codigo, self.carrera.nombre, 18, "Municipal", self.ces.nombre, "", "A"])

        buffer = BytesIO()
        workbook.save(buffer)
        file = SimpleUploadedFile("plan-plazas-province.xlsx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self._import_plan_plaza(file)

        self.assertEqual(response.status_code, 201, response.content)
        plan = self.carrera.plan_plaza.order_by("-id").first()
        self.assertIsNotNone(plan)
        self.assertEqual(plan.provincia_id, self.provincia.id)
        self.assertEqual(plan.proceso.etapa_id, plan_stage.id)
        self.assertEqual(plan.proceso.anio.year, 2026)

    def test_import_resolves_national_career_code_for_user_province(self):
        from io import BytesIO
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile

        other_province = Provincia.objects.create(nombre="Santiago de Cuba")
        self.user.provincia = other_province
        self.user.save(update_fields=["provincia"])

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plan de Plazas"
        sheet.append(["Código_Carrera", "Nombre_Carrera", "Cantidad_Plazas", "Tipo_Otorgamiento", "CES", "Provincia", "Sexo"])
        sheet.append([self.carrera.codigo, self.carrera.nombre, 10, "Municipal", self.ces.nombre, other_province.nombre, "A"])

        buffer = BytesIO()
        workbook.save(buffer)
        file = SimpleUploadedFile("plan-plazas-national-code.xlsx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self._import_plan_plaza(file)

        self.assertEqual(response.status_code, 201, response.content)
        plan = PlanPlaza.objects.get(carrera=self.carrera, provincia=other_province)
        self.assertEqual(plan.carrera_id, self.carrera.id)
        self.assertEqual(plan.provincia_id, other_province.id)

    def test_import_rejects_invalid_type_values(self):
        from io import BytesIO
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plan de Plazas"
        sheet.append(["Código_Carrera", "Nombre_Carrera", "Cantidad_Plazas", "Tipo_Otorgamiento", "CES", "Provincia", "Sexo"])
        sheet.append([self.carrera.codigo, self.carrera.nombre, 5, "Preuniversitario", self.ces.nombre, self.provincia.nombre, "F"])

        buffer = BytesIO()
        workbook.save(buffer)
        file = SimpleUploadedFile("plan-plazas-invalid.xlsx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self._import_plan_plaza(file)

        self.assertEqual(response.status_code, 400)
        self.assertIn("Tipo_Otorgamiento", str(response.data))

    def test_list_remains_empty_when_no_current_year_process_exists(self):
        self.user.provincia = self.provincia
        self.user.save(update_fields=["provincia"])

        plan_stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[3])
        Proceso.objects.filter(etapa=plan_stage, anio__year=timezone.now().year).delete()

        response = self.client.get(reverse("plan-plazas"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, [])

    def test_landing_returns_active_year_and_user_province_plan(self):
        self.user.provincia = self.provincia
        self.user.save(update_fields=["provincia"])

        plan_stage = Etapa.objects.get(nombre=ETAPAS_NOMBRES[3])
        plan_stage.estado = "en_curso"
        plan_stage.save(update_fields=["estado"])

        active_process = Proceso.objects.create(anio=date(timezone.now().year, 1, 1), etapa=plan_stage)
        PlanPlaza.objects.create(
            proceso=active_process,
            carrera=self.carrera,
            cantidad_plazas=20,
            otorgamiento_tipo=TipoOtorgamiento.objects.get(nombre="Municipal"),
            ces=self.ces,
            provincia=self.provincia,
            sexo="A",
        )

        response = self.client.get(reverse("plan-plazas-landing"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["year"], timezone.now().year)
        self.assertEqual(response.data["provincia_nombre"], self.provincia.nombre)
        self.assertEqual(len(response.data["items"]), 1)
        self.assertEqual(response.data["items"][0]["cantidad_plazas"], 20)


class CareerInformationLandingTests(APITestCase):
    def setUp(self):
        self.holguin = Provincia.objects.create(nombre="Holguín")
        self.other_province = Provincia.objects.create(nombre="Matanzas")
        self.ces = Ces.objects.create(nombre="CES catálogo landing")
        self.holguin_career = Carrera.objects.create(
            codigo="LAND-HOL-1",
            nombre="Carrera informada de Holguín",
            descripcion="Ficha pública disponible.",
            ces=self.ces,
            provincia=self.holguin,
        )
        Carrera.objects.create(
            codigo="LAND-HOL-2",
            nombre="Carrera sin ficha",
            ces=self.ces,
            provincia=self.holguin,
        )
        self.other_career = Carrera.objects.create(
            codigo="LAND-MAT-1",
            nombre="Carrera informada de Matanzas",
            campo_laboral="Empresas del sector.",
            ces=self.ces,
            provincia=self.other_province,
        )

    def test_anonymous_landing_defaults_to_holguin_and_hides_empty_records(self):
        response = self.client.get(reverse("career-information-landing"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["provincia_nombre"], "Holguín")
        self.assertEqual([item["id"] for item in response.data["items"]], [self.holguin_career.id])

    def test_authenticated_user_sees_informed_careers_from_own_province(self):
        user = Usuario.objects.create_user(
            username="landing.provincial",
            email="landing.provincial@example.com",
            password="secret1234",
            rol="jefe_comision",
            provincia=self.other_province,
        )
        self.client.force_authenticate(user)

        response = self.client.get(reverse("career-information-landing"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["provincia_nombre"], "Matanzas")
        self.assertEqual([item["id"] for item in response.data["items"]], [self.other_career.id])


class ProvincialCareerInformationTests(APITestCase):
    def setUp(self):
        self.province = Provincia.objects.create(nombre="Provincia carreras")
        self.ces = Ces.objects.create(nombre="CES carreras")
        user = Usuario.objects.create_user(
            username="jefe.carreras",
            email="jefe.carreras@example.com",
            password="secret1234",
            rol="jefe_comision",
            provincia=self.province,
        )
        self.client.force_authenticate(user)

    def test_create_and_edit_career_information(self):
        fields = {
            "codigo": "INFO-1",
            "nombre": "Carrera con perfil",
            "descripcion": "Descripción",
            "campo_laboral": "Campo laboral",
            "requisitos": "Requisitos",
            "perfil_egresado": "Perfil del egresado",
            "ces": self.ces.id,
            "provincia": self.province.id,
        }
        created = self.client.post(reverse("provincial-carreras"), fields, format="json")

        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["descripcion"], "Descripción")
        self.assertEqual(created.data["campo_laboral"], "Campo laboral")
        self.assertEqual(created.data["requisitos"], "Requisitos")
        self.assertEqual(created.data["perfil_egresado"], "Perfil del egresado")

        updated = self.client.patch(
            reverse("provincial-carrera-detail", args=[created.data["id"]]),
            {"campo_laboral": "Campo laboral actualizado"},
            format="json",
        )

        self.assertEqual(updated.status_code, 200, updated.data)
        self.assertEqual(updated.data["campo_laboral"], "Campo laboral actualizado")
        self.assertEqual(updated.data["requisitos"], "Requisitos")


class JefeStageReadOnlyTests(APITestCase):
    """Con la etapa 3 completada el panel del jefe es de solo lectura: GET ok, escrituras 403."""

    def setUp(self):
        self.provincia = Provincia.objects.create(nombre="Provincia RO")
        self.ces = Ces.objects.create(nombre="CES RO", activa=True)
        self.carrera = Carrera.objects.create(codigo="RO-1", nombre="Carrera RO", ces=self.ces, provincia=self.provincia, activa=True)
        self.tipo = TipoOtorgamiento.objects.create(nombre="Municipal RO", activa=True)
        self.jefe = Usuario.objects.create_user(
            username="jefe.ro", email="jefe.ro@example.com", password="secret1234",
            rol="jefe_comision", provincia=self.provincia,
        )
        self.client.force_authenticate(self.jefe)
        self.etapa = Etapa.objects.get(nombre=ETAPAS_NOMBRES[3])
        self.etapa.estado = "completada"
        self.etapa.save(update_fields=["estado"])
        self.proceso = Proceso.objects.get_or_create(anio=timezone.now().date().replace(month=1, day=1), etapa=self.etapa)[0]
        self.plan = PlanPlaza.objects.create(
            proceso=self.proceso, carrera=self.carrera, ces=self.ces, provincia=self.provincia,
            otorgamiento_tipo=self.tipo, cantidad_plazas=5, sexo="A",
        )

    def _payload(self):
        return {
            "proceso": self.proceso.id, "carrera": self.carrera.id, "ces": self.ces.id,
            "provincia": self.provincia.id, "otorgamiento_tipo": self.tipo.id, "cantidad_plazas": 3, "sexo": "A",
        }

    def test_get_still_returns_data_when_stage_completed(self):
        response = self.client.get(reverse("plan-plazas"))
        self.assertEqual(response.status_code, 200, response.content)

    def test_crud_rejected_when_stage_completed(self):
        base = reverse("plan-plazas")
        detail = reverse("plan-plaza-detail", args=[self.plan.id])
        self.assertEqual(self.client.post(base, self._payload(), format="json").status_code, 403)
        self.assertEqual(self.client.patch(detail, {"cantidad_plazas": 9}, format="json").status_code, 403)
        self.assertEqual(self.client.delete(detail).status_code, 403)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.cantidad_plazas, 5)

    def test_crud_allowed_when_stage_in_progress(self):
        self.etapa.estado = "en_curso"
        self.etapa.save(update_fields=["estado"])
        response = self.client.patch(reverse("plan-plaza-detail", args=[self.plan.id]), {"cantidad_plazas": 9}, format="json")
        self.assertEqual(response.status_code, 200, response.content)

    def test_import_plan_plaza_rejected_when_stage_completed(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        file = SimpleUploadedFile("p.xlsx", b"x", content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response = self.client.post(reverse("import-plan-plaza"), {"file": file}, format="multipart")
        self.assertEqual(response.status_code, 403)
